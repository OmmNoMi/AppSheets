#!/usr/bin/env python3
"""
build_cini_flat_survey_engine.py
================================
Builds the complete "One Form, One Table" (CINI Pattern) survey engine:
1. Single flat 'Survey' table containing all questionnaire columns.
2. Trilingual 'AppVariables' dictionary (English, Hindi, Rajasthani).
3. 'AppUser' table with Language Switcher & Daily Target tracker.
4. 'SamplingFrame' master table for pilot districts.
5. Formatted Excel Master Workbooks with OmmNoMi brand styling.
"""

import os
import sys
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

PROJECT_DIR = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs"
DATA_DIR = os.path.join(PROJECT_DIR, "data")
os.makedirs(DATA_DIR, exist_ok=True)
NOW_STR = datetime.now().strftime("%m/%d/%Y %H:%M:%S")

STANDARD_APPVAR_HEADERS = [
    'ID', 'Table', 'Column', 'Tags', 'ValueControl', 'Title', 'Description', 'UsedFor',
    'Decimal', 'EnumValue', 'EnumList', 'VariableList', 'DateValue', 'Photo', 'URL', 'File',
    'Title_hi', 'Title_raj', 'ActionIcon', 'LastEditBy', 'LastEditOn'
]

app_vars = []

def add_var(id_val, table='Survey', col='', tags='', vc='Enum', title='', desc='', used_for='',
            enum_val='', var_list='', enum_list='', title_hi='', title_raj='', icon='', decimal=None):
    row = {h: '' for h in STANDARD_APPVAR_HEADERS}
    row['ID'] = str(id_val).strip()
    row['Table'] = str(table).strip()
    row['Column'] = str(col).strip()
    row['Tags'] = str(tags).strip()
    row['ValueControl'] = str(vc).strip()
    row['Title'] = str(title).strip()
    row['Description'] = str(desc).strip()
    row['UsedFor'] = str(used_for).strip()
    row['EnumValue'] = str(enum_val if enum_val else title).strip()
    row['VariableList'] = str(var_list).strip()
    row['EnumList'] = str(enum_list).strip()
    row['Title_hi'] = str(title_hi).strip()
    row['Title_raj'] = str(title_raj).strip()
    row['ActionIcon'] = str(icon).strip()
    if decimal is not None:
        row['Decimal'] = str(decimal)
    row['LastEditBy'] = 'Antigravity'
    row['LastEditOn'] = NOW_STR
    app_vars.append(row)

# ==========================================
# 1. SYSTEM & CONFIGURATION
# ==========================================
add_var('CompanyName', 'AppVariables', '', 'System, Config', 'Enum', 'Company Name', 'CmF & RAJEEVIKA Study', 'Header in reports', 'Centre for microFinance (CmF)', title_hi='सेंटर फॉर माइक्रोफाइनेंस (CmF)', title_raj='सेंटर फॉर माइक्रोफाइनेंस (CmF)')
add_var('AppName', 'AppVariables', '', 'System, Config', 'Enum', 'App Name', 'SHG Women Entrepreneurs Survey', 'Application Title', 'SHG Women Entrepreneurs Study', title_hi='एसएचजी महिला उद्यमी अध्ययन', title_raj='एसएचजी महिला उद्यमी अध्ययन')

# Languages supported
languages = [
    ('LANG_EN', 'English', 'अंग्रेज़ी', 'अंग्रेजी'),
    ('LANG_HI', 'Hindi', 'हिन्दी', 'हिंदी'),
    ('LANG_RAJ', 'Rajasthani', 'राजस्थानी', 'राजस्थानी')
]
for lid, en, hi, raj in languages:
    add_var(lid, 'AppVariables', 'Language', 'Option, System', 'Enum', en, '', 'Language Option', en, title_hi=hi, title_raj=raj)
add_var('Q_SYS_LANG', 'AppUser', 'Language', 'QuestionPrompt, System', 'VariableList', 'Preferred Language', 'User display language', var_list='LANG_EN , LANG_HI , LANG_RAJ', title_hi='पसंदीदा भाषा', title_raj='पसंदीदा भासा')

# User Roles (SOP §3)
user_roles = [
    ('ROLE_INVESTIGATOR', 'Field Investigator', 'फील्ड अन्वेषक', 'फील्ड अन्वेषक', 'user'),
    ('ROLE_SUPERVISOR', 'Field Supervisor', 'फील्ड सुपरवाइजर', 'फील्ड सुपरवाइजर', 'shield'),
    ('ROLE_ADMIN', 'Admin', 'एडमिन', 'एडमिन', 'settings')
]
for rid, en, hi, raj, icon in user_roles:
    add_var(rid, 'AppVariables', 'Role', 'Option, System', 'Enum', en, '', 'Role Option', en, title_hi=hi, title_raj=raj, icon=icon)
add_var('Q_SYS_ROLE', 'AppUser', 'Role', 'QuestionPrompt, System', 'VariableList', 'User Role', 'AppUser role selector', var_list='ROLE_INVESTIGATOR , ROLE_SUPERVISOR , ROLE_ADMIN', title_hi='उपयोगकर्ता भूमिका', title_raj='उपयोगकर्ता भूमिका')

# Survey Statuses
statuses = [
    ('STAT_DRAFT', 'Draft', 'ड्राफ्ट / अधूरा', 'कच्चो / अधूरो'),
    ('STAT_SUBMITTED', 'Submitted', 'पूर्ण / जमा किया गया', 'पूरो / जमा कर दियो'),
    ('STAT_VERIFIED', 'Verified', 'सत्यापित', 'जांचियो गयो')
]
for sid, en, hi, raj in statuses:
    add_var(sid, 'AppVariables', 'Status', 'Option, System', 'Enum', en, '', 'Status Option', en, title_hi=hi, title_raj=raj)

# Section Completion Statuses (Done / Not Done)
sec_statuses = [
    ('SEC_NOT_STARTED', 'Not Started', 'शुरू नहीं हुआ', 'सरू कोनी हुयो'),
    ('SEC_IN_PROGRESS', 'In Progress', 'जारी है / अधूरा', 'चालू छै / अधूरो'),
    ('SEC_DONE', 'Done', 'पूर्ण (Done)', 'पूरो (Done)'),
    ('SEC_NA', 'Not Applicable', 'लागू नहीं (N/A)', 'लागू कोनी (N/A)')
]
sec_stat_ids = [s[0] for s in sec_statuses]
for sid, en, hi, raj in sec_statuses:
    add_var(sid, 'AppVariables', 'SectionStatus', 'Option, System', 'Enum', en, '', 'Section Status Option', en, title_hi=hi, title_raj=raj)

# Section Status Prompts on Survey Table
add_var('Q_STAT_PROFILE', 'Survey', 'Status_Profile', 'QuestionPrompt, SectionStatus', 'Enum', 'Profile Section Status', 'Section B Status', 'Status of Profile section', 'Not Started', var_list=' , '.join(sec_stat_ids), title_hi='प्रोफ़ाइल भाग स्थिति', title_raj='प्रोफ़ाइल भाग स्थिति')
add_var('Q_STAT_OPERATIONS', 'Survey', 'Status_Operations', 'QuestionPrompt, SectionStatus', 'Enum', 'Operations Section Status', 'Section C Status', 'Status of Operations section', 'Not Started', var_list=' , '.join(sec_stat_ids), title_hi='उद्यम संचालन भाग स्थिति', title_raj='उद्यम संचालन भाग स्थिति')
add_var('Q_STAT_CHALLENGES', 'Survey', 'Status_Challenges', 'QuestionPrompt, SectionStatus', 'Enum', 'Challenges Section Status', 'Section D Status', 'Status of Challenges section', 'Not Started', var_list=' , '.join(sec_stat_ids), title_hi='चुनौतियां भाग स्थिति', title_raj='अड़चनां भाग स्थिति')
add_var('Q_STAT_SCHEME', 'Survey', 'Status_SchemeImpact', 'QuestionPrompt, SectionStatus', 'Enum', 'Scheme Impact Section Status', 'Section E Status', 'Status of Scheme Impact section', 'Not Started', var_list=' , '.join(sec_stat_ids), title_hi='योजना प्रभाव भाग स्थिति', title_raj='योजना प्रभाव भाग स्थिति')
add_var('Q_STAT_DIGITAL', 'Survey', 'Status_Digital', 'QuestionPrompt, SectionStatus', 'Enum', 'Digital Media Section Status', 'Section F Status', 'Status of Digital Media section', 'Not Started', var_list=' , '.join(sec_stat_ids), title_hi='डिजिटल मीडिया भाग स्थिति', title_raj='डिजिटल मीडिया भाग स्थिति')
add_var('Q_STAT_POST_EXIT', 'Survey', 'Status_PostExit', 'QuestionPrompt, SectionStatus', 'Enum', 'Post-Exit OSF Section Status', 'Section G Status', 'Status of Post-Exit OSF section', 'Not Started', var_list=' , '.join(sec_stat_ids), title_hi='पोस्ट-एग्जिट OSF भाग स्थिति', title_raj='पोस्ट-एग्जिट OSF भाग स्थिति')

# Action Grid Navigation Buttons (Displayed prominently at top of Survey Detail View)
action_grid_buttons = [
    ('ACT_Complete', 'Complete', 'पूर्ण', 'पूरो', 'check-circle', 'Submit & lock completed survey'),
    ('ACT_Profile', 'Profile', 'प्रोफ़ाइल', 'प्रोफ़ाइल', 'user', 'Navigate to Section B: Respondent & Household Profile'),
    ('ACT_Operations', 'Operations', 'उद्यम संचालन', 'काम-धंधो', 'storefront', 'Navigate to Section C: Enterprise Operations & Capital'),
    ('ACT_Challenges', 'Challenges', 'चुनौतियां', 'अड़चनां', 'trending-up', 'Navigate to Section D: Ease of Doing Business & Challenges'),
    ('ACT_SchemeImpact', 'Scheme Impact', 'योजना प्रभाव', 'लोन रो फायदो', 'account-balance', 'Navigate to Section E: Impact of SVEP/OSF Schemes'),
    ('ACT_Digital', 'Digital & Media', 'डिजिटल एवं मीडिया', 'ऑनलाइन लेन-देन', 'smartphone', 'Navigate to Section F: Online Transactions & Social Media'),
    ('ACT_PostExit', 'Post-Exit OSF', 'पोस्ट-एग्जिट OSF', 'पुरानी स्थिति', 'history', 'Navigate to Section G: Post-Exit OSF in Baran & Ratangarh')
]
for aid, en, hi, raj, icon, desc in action_grid_buttons:
    add_var(aid, 'AppVariables', 'ActionGrid', 'UI, ActionGrid, Navigation', 'Enum', en, desc, 'Action Grid button', en, title_hi=hi, title_raj=raj, icon=icon)

# Yes / No Standard
add_var('OPT_YES', 'AppVariables', 'YesNo', 'Option', 'Enum', 'Yes', '', 'Yes Option', 'Yes', title_hi='हाँ', title_raj='हाँ')
add_var('OPT_NO', 'AppVariables', 'YesNo', 'Option', 'Enum', 'No', '', 'No Option', 'No', title_hi='नहीं', title_raj='कोनी / ना')

# ==========================================
# 2. GEOGRAPHY (DISTRICTS & BLOCKS)
# ==========================================
districts = [
    ('DIST_BARAN', 'Baran', 'बारां', 'बारां', 'Fertile East'),
    ('DIST_CHURU', 'Churu', 'चूरू', 'चूरू', 'Arid West'),
    ('DIST_DAUSA', 'Dausa', 'दौसा', 'दौसा', 'Semi-Arid'),
    ('DIST_DUNGARPUR', 'Dungarpur', 'डूंगरपुर', 'डूंगरपुर', 'Southern Tribal'),
    ('DIST_JODHPUR', 'Jodhpur', 'जोधपुर', 'जोधपुर', 'Western Arid')
]
dist_ids = []
for did, en, hi, raj, zone in districts:
    add_var(did, 'AppVariables', 'District', 'Option, Geography', 'Enum', en, zone, 'District Option', en, title_hi=hi, title_raj=raj)
    dist_ids.append(did)

add_var('Q_A_01_00', 'Survey', 'District', 'QuestionPrompt, SectionA, Order:1', 'VariableList', 'District', 'Section A Q1', 'District selector', var_list=' , '.join(dist_ids), title_hi='जिला', title_raj='जिलो', decimal=1)

# State HQ for System Admins (AppUser)
add_var('DIST_JAIPUR', 'AppVariables', 'District', 'Option, Geography, Admin', 'Enum', 'Jaipur (HQ)', 'State Headquarter', 'District Option', 'Jaipur (HQ)', title_hi='जयपुर (मुख्यालय)', title_raj='जयपुर (मुख्यालय)')

blocks = [
    # Baran
    ('BLK_CHHIPABAROD', 'Chhipabarod', 'छीपाबड़ौद', 'छीपाबड़ौद', 'DIST_BARAN'),
    ('BLK_BARAN', 'Baran', 'बारां', 'बारां', 'DIST_BARAN'),
    # Churu
    ('BLK_RATANGARH', 'Ratangarh', 'रतनगढ़', 'रतनगढ़', 'DIST_CHURU'),
    ('BLK_SUJANGARH', 'Sujangarh', 'सुजानगढ़', 'सुजानगढ़', 'DIST_CHURU'),
    # Dausa
    ('BLK_SIKANDRA', 'Sikandra', 'सिकंदरा', 'सिकंदरा', 'DIST_DAUSA'),
    # Dungarpur
    ('BLK_SAGWARA', 'Sagwara', 'सागवाड़ा', 'सागवाड़ा', 'DIST_DUNGARPUR'),
    ('BLK_GALIAKOT', 'Galiakot', 'गलियाकोट', 'गलियाकोट', 'DIST_DUNGARPUR'),
    # Jodhpur
    ('BLK_MANDOR', 'Mandor', 'मंडोर', 'मंडोर', 'DIST_JODHPUR'),
    ('BLK_LUNI', 'Luni', 'लूणी', 'लूणी', 'DIST_JODHPUR'),
    ('BLK_SHERGADH', 'Shergadh', 'शेरगढ़', 'शेरगढ़', 'DIST_JODHPUR')
]
blk_ids = []
for bid, en, hi, raj, parent_dist in blocks:
    add_var(bid, 'AppVariables', 'Block', f'Option, Geography, Parent:{parent_dist}', 'Enum', en, parent_dist, 'Block Option', en, title_hi=hi, title_raj=raj)
    blk_ids.append(bid)

add_var('Q_A_02_00', 'Survey', 'Block', 'QuestionPrompt, SectionA, Order:2', 'VariableList', 'Block', 'Section A Q2', 'Block selector', var_list=' , '.join(blk_ids), title_hi='ब्लॉक / खंड', title_raj='ब्लॉक', decimal=2)

# Section A Basic Details Prompts
add_var('Q_A_03_00', 'Survey', 'VillageGP', 'QuestionPrompt, SectionA, Order:3', 'Text', 'Village / Gram Panchayat', 'Section A Q3', 'Village name', title_hi='गांव / ग्राम पंचायत', title_raj='गांम / ग्राम पंचायत', decimal=3)
add_var('Q_A_04_00', 'Survey', 'RespondentName', 'QuestionPrompt, SectionA, Order:4', 'Name', 'Respondent Name', 'Section A Q4', 'Woman entrepreneur name', title_hi='उत्तरदाता का नाम', title_raj='उत्तरदाता रो नाम', decimal=4)
add_var('Q_A_04_01', 'Survey', 'ContactNumber', 'QuestionPrompt, SectionA, Order:5', 'Phone', 'Contact Number', 'Section A Q4.1', 'Phone number', title_hi='संपर्क नंबर (मोबाइल)', title_raj='मोबाइल नंबर', decimal=5)
add_var('Q_A_05_00', 'Survey', 'SHGName', 'QuestionPrompt, SectionA, Order:6', 'Text', 'SHG Name', 'Section A Q5', 'Self Help Group', title_hi='स्वयं सहायता समूह (SHG) का नाम', title_raj='एसएचजी रो नाम', decimal=6)
add_var('Q_A_06_00', 'Survey', 'VOName', 'QuestionPrompt, SectionA, Order:7', 'Text', 'VO Name', 'Section A Q6', 'Village Organization', title_hi='ग्राम संगठन (VO) का नाम', title_raj='ग्राम संगठन (VO) रो नाम', decimal=7)
add_var('Q_A_07_00', 'Survey', 'CLFName', 'QuestionPrompt, SectionA, Order:8', 'Text', 'CLF Name', 'Section A Q7', 'Cluster Level Federation', title_hi='सीएलएफ का नाम', title_raj='सीएलएफ रो नाम', decimal=8)
add_var('Q_A_08_00', 'Survey', 'SHGMembershipYears', 'QuestionPrompt, SectionA, Order:9', 'Number', 'Years of SHG membership', 'Section A Q8', 'SHG membership duration', title_hi='SHG सदस्यता के वर्ष', title_raj='समूह में जुड़ाव रा साल', decimal=9)
add_var('Q_A_09_00', 'Survey', 'LeadershipRole', 'QuestionPrompt, SectionA, Order:10', 'Enum', 'Have you been in a leadership role in SHG/CLF/VO?', 'Section A Q9', 'Leadership status', var_list='OPT_YES , OPT_NO', title_hi='क्या आप SHG/CLF/VO में किसी नेतृत्व पद पर रही हैं?', title_raj='कांई थे समूह/VO/CLF में पदाधिकारी या अगवाण रह्या छो?', decimal=10)
add_var('Q_A_10_00', 'Survey', 'LeadershipYears', 'QuestionPrompt, SectionA, Order:11', 'Number', 'Years of experience in leadership roles?', 'Section A Q10', 'Leadership tenure', title_hi='नेतृत्व पदों पर अनुभव के वर्ष', title_raj='पदाधिकारी के रूप में कित्ता साल रो अनुभव छै?', decimal=11)
add_var('Q_A_11_00', 'Survey', 'RelatedToCRP', 'QuestionPrompt, SectionA, Order:12', 'Enum', 'Are you related to any of the SVEP/OSF CRP?', 'Section A Q11', 'Related to CRP', var_list='OPT_YES , OPT_NO', title_hi='क्या आप किसी SVEP/OSF सीआरपी से संबंधित/रिश्तेदार हैं?', title_raj='कांई थे कोई SVEP/OSF सीआरपी रा सगा-संबंधी छो?', decimal=12)

# Intervention types
interventions = [
    ('INT_SVEP', 'SVEP', 'SVEP', 'SVEP'),
    ('INT_OSF', 'OSF', 'OSF', 'OSF'),
    ('INT_OSF_PHASED', 'OSF phased out', 'OSF समाप्त (Phased out)', 'OSF बंद होग्यो'),
    ('INT_DONT_KNOW', "Don't know", 'पता नहीं', 'ठा कोनी')
]
int_ids = []
for iid, en, hi, raj in interventions:
    add_var(iid, 'AppVariables', 'InterventionType', 'Option', 'Enum', en, '', 'Intervention Option', en, title_hi=hi, title_raj=raj)
    int_ids.append(iid)

add_var('Q_A_12_00', 'Survey', 'EPInterventionType', 'QuestionPrompt, SectionA, Order:13', 'Enum', 'Type of Enterprise Promotion (EP) intervention', 'Section A Q12', 'EP Scheme type', var_list=' , '.join(int_ids), title_hi='उद्यम संवर्धन (EP) हस्तक्षेप का प्रकार', title_raj='उद्यम सहायता रो प्रकार', decimal=13)
add_var('Q_A_13_00', 'Survey', 'EnterpriseName', 'QuestionPrompt, SectionA, Order:14', 'Text', 'Enterprise Name', 'Section A Q13', 'Name of primary enterprise', title_hi='उद्यम / दुकान का नाम', title_raj='दुकान / काम-धंधे रो नाम', decimal=14)
add_var('Q_A_13_01', 'Survey', 'ParallelEnterpriseName', 'QuestionPrompt, SectionA, Order:15', 'Text', 'Parallel Enterprise Name (if running two businesses)', 'Section A Q13.1', 'Second enterprise name', title_hi='दूसरे समानांतर उद्यम का नाम (यदि दो व्यवसाय हैं)', title_raj='दूजे काम-धंधे रो नाम (अगर दो चलावै छो)', decimal=15)
add_var('Q_A_14_00', 'Survey', 'EnterpriseSetupYear', 'QuestionPrompt, SectionA, Order:16', 'Number', 'Years of setting up enterprise (Year or count)', 'Section A Q14', 'Enterprise setup year', title_hi='उद्यम स्थापित करने का वर्ष / अवधि', title_raj='काम-धंधो कदै सुरू करयो छो (साल)?', decimal=16)
add_var('Q_A_15_00', 'Survey', 'LoanReceivedYear', 'QuestionPrompt, SectionA, Order:17', 'Number', 'Year of receiving SVEP/OSF loan?', 'Section A Q15', 'Year loan received', title_hi='SVEP/OSF ऋण प्राप्त करने का वर्ष', title_raj='योजना रो लोन किस साल मिल्यो छो?', decimal=17)

# Business Type
biz_types = [
    ('BTY_TRADING', 'Trading', 'व्यापार / ट्रेडिंग', 'खरीद-बिक्री (ट्रेडिंग)'),
    ('BTY_SERVICING', 'Servicing', 'सेवा / सर्विसिंग', 'सेवा (सर्विसिंग)'),
    ('BTY_MANUFACTURING', 'Manufacturing / production', 'उत्पादन / विनिर्माण', 'बणावण रो काम (उत्पादन)')
]
bty_ids = []
for bid, en, hi, raj in biz_types:
    add_var(bid, 'AppVariables', 'BusinessType', 'Option', 'Enum', en, '', 'Business Type Option', en, title_hi=hi, title_raj=raj)
    bty_ids.append(bid)
add_var('Q_A_16_00', 'Survey', 'BusinessType', 'QuestionPrompt, SectionA, Order:18', 'EnumList', 'Type of business enterprise (Multiselect)', 'Section A Q16', 'Category of enterprise', var_list=' , '.join(bty_ids), title_hi='व्यावसायिक उद्यम का प्रकार', title_raj='काम-धंधे रो मुख्य प्रकार', decimal=18)

# Business Activities
activities = [
    # Trading
    ('ACT_VEG_FRUIT', 'Vegetable / Fruit', 'सब्जी / फल', 'सब्जी / फल', 'Trading'),
    ('ACT_GROCERY', 'Grocery', 'किराना', 'किराणा', 'Trading'),
    ('ACT_FANCY_STORE', 'Fancy / Cosmetic / General store', 'फैंसी / कॉस्मेटिक / जनरल स्टोर', 'प्रसाधन / जनरल स्टोर', 'Trading'),
    ('ACT_APPAREL', 'Apparel / fabric', 'कपड़ा / रेडीमेड वस्त्र', 'कपड़ा / वेशभूषा', 'Trading'),
    ('ACT_ELECTRIC_GOODS', 'Electric goods', 'इलेक्ट्रिक सामान', 'बिजली रो सामान', 'Trading'),
    ('ACT_STONE_SHOP', 'Stone shop', 'पत्थर की दुकान', 'भाटां/पत्थर री दुकान', 'Trading'),
    # Servicing
    ('ACT_FLOUR_MILL', 'Flour mill (Chakki)', 'आटा चक्की', 'आटा चक्की', 'Servicing'),
    ('ACT_TAILORING', 'Tailoring', 'सिलाई / टेलरिंग', 'सिलाई / दर्जी काम', 'Servicing'),
    ('ACT_BEAUTY_PARLOUR', 'Beauty parlour', 'ब्यूटी पार्लर', 'ब्यूटी पार्लर', 'Servicing'),
    ('ACT_AUTO_REPAIR', 'Auto-mechanic / two-wheeler repair', 'ऑटो मैकेनिक / दोपहिया मरम्मत', 'गाड़ी/मोटरसाइकिल मिस्त्री', 'Servicing'),
    ('ACT_EMITRA', 'E-mitra / Online kiosk', 'ई-मित्र / ऑनलाइन केंद्र', 'ई-मित्र केंद्र', 'Servicing'),
    ('ACT_TRANSPORT', 'Transport', 'परिवहन / वाहन', 'गाड़ी भाड़ा / परिवहन', 'Servicing'),
    ('ACT_TENT_HOUSE', 'Tent house', 'टेंट हाउस', 'टेंट हाउस', 'Servicing'),
    ('ACT_MOBILE_REPAIR', 'Mobile repair shop', 'मोबाइल रिपेयर दुकान', 'मोबाइल ठीक करण री दुकान', 'Servicing'),
    ('ACT_STONE_CUTTING', 'Stone cutting', 'पत्थर कटाई', 'पत्थर कटाई', 'Servicing'),
    # Production
    ('ACT_SANITARY_NAPKIN', 'Sanitary napkin making', 'सैनिटरी नैपकिन निर्माण', 'सैनिटरी पैड बणावण', 'Production'),
    ('ACT_HANDICRAFT', 'Handicraft', 'हस्तशिल्प / कसीदाकारी', 'हाथ रो काम / कसीदाकारी', 'Production'),
    ('ACT_DAIRY_MILK', 'Dairy shop / Milk collection centre', 'डेयरी दुकान / दुग्ध संकलन केंद्र', 'दूध डेरी / संकलन केंद्र', 'Production'),
    ('ACT_JUICE', 'Juice', 'जूस की दुकान', 'जूस री दुकान', 'Production'),
    ('ACT_FOOD_PROCESSING', 'Food processing (pickle/badi/papad)', 'खाद्य प्रसंस्करण (अचार/बड़ी/पापड़ निर्माण)', 'अचार, पापड़, बड़ी बणावण', 'Production'),
    ('ACT_FOOD_MAKING', 'Food making (Sweets/Namkeen/hotel)', 'मिठाई / नमकीन / ढाबा / होटल', 'मिठाई / नमकीन / होटल', 'Production'),
    ('ACT_SWEET_BOX', 'Sweet box making', 'मिठाई के डिब्बे बनाना', 'मिठाई रा डिब्बा बणावण', 'Production'),
    ('ACT_FLAG_MAKING', 'Flag making', 'झंडा निर्माण', 'झंडा बणावण', 'Production'),
    ('ACT_LEATHER_PRODUCTS', 'Leather products', 'चमड़े के उत्पाद / जूते', 'चामड़ा रो काम / जूता', 'Production'),
    ('ACT_STONE_IDOLS', 'Stone idols', 'पत्थर की मूर्तियां', 'पत्थर री मूर्तियां बणावण', 'Production'),
    ('ACT_ANY_OTHER', 'Any other activity', 'अन्य कोई गतिविधि', 'दूजो कोई काम', 'Other')
]
act_ids = []
for aid, en, hi, raj, grp in activities:
    add_var(aid, 'AppVariables', 'Activity', f'Option, Trade, Group:{grp}', 'Enum', en, grp, 'Activity Option', en, title_hi=hi, title_raj=raj)
    act_ids.append(aid)
add_var('Q_A_17_00', 'Survey', 'BusinessActivities', 'QuestionPrompt, SectionA, Order:19', 'EnumList', 'Main business activities of the enterprise (Multiselect)', 'Section A Q17', 'Specific trade activities', var_list=' , '.join(act_ids), title_hi='उद्यम की मुख्य व्यावसायिक गतिविधियां', title_raj='काम-धंधे री मुख्य गतिविधियां', decimal=19)
add_var('Q_A_17_01', 'Survey', 'BusinessActivitiesOther', 'QuestionPrompt, SectionA, Order:20', 'Text', 'Specify other business activity (if selected Any other)', 'Section A Q17.1', 'Other trade specification', title_hi='अन्य व्यावसायिक गतिविधि का विवरण दें', title_raj='दूजे काम रो ब्योरो लिखो', decimal=20)

# ==========================================
# 3. SECTION B: RESPONDENT & HOUSEHOLD PROFILE
# ==========================================
age_cohorts = [
    ('AGE_18_25', '18-25', '18-25 वर्ष', '18-25 साल'),
    ('AGE_26_35', '26-35', '26-35 वर्ष', '26-35 साल'),
    ('AGE_36_45', '36-45', '36-45 वर्ष', '36-45 साल'),
    ('AGE_46_55', '46-55', '46-55 वर्ष', '46-55 साल'),
    ('AGE_ABOVE_55', 'Above 55', '55 वर्ष से अधिक', '55 साल सूं बत्ता')
]
age_ids = [a[0] for a in age_cohorts]
for aid, en, hi, raj in age_cohorts:
    add_var(aid, 'AppVariables', 'AgeCohort', 'Option', 'Enum', en, '', 'Age Cohort Option', en, title_hi=hi, title_raj=raj)
add_var('Q_B_01_00', 'Survey', 'RespondentAge', 'QuestionPrompt, SectionB, Order:21', 'Enum', 'What is the age of the respondent?', 'Section B Q1', 'Respondent age', var_list=' , '.join(age_ids), title_hi='उत्तरदाता की आयु क्या है?', title_raj='थारी उमर कितनी छै?', decimal=21)

maritals = [
    ('MAR_SINGLE', 'Single', 'अविवाहित', 'अणपरणी / कुंवारी'),
    ('MAR_MARRIED', 'Married', 'विवाहित', 'परणी'),
    ('MAR_WIDOWED', 'Widowed', 'विधवा', 'रांड/विधवा'),
    ('MAR_SEPARATED', 'Separated', 'परित्यक्ता / अलग', 'अलग रहवै'),
    ('MAR_DIVORCED', 'Divorced', 'तलाकशुदा', 'तलाकशुदा')
]
mar_ids = [m[0] for m in maritals]
for mid, en, hi, raj in maritals:
    add_var(mid, 'AppVariables', 'MaritalStatus', 'Option', 'Enum', en, '', 'Marital Option', en, title_hi=hi, title_raj=raj)
add_var('Q_B_02_00', 'Survey', 'MaritalStatus', 'QuestionPrompt, SectionB, Order:22', 'Enum', 'What is the marital status?', 'Section B Q2', 'Marital status', var_list=' , '.join(mar_ids), title_hi='वैवाहिक स्थिति क्या है?', title_raj='वैवाहिक स्थिति कांई छै?', decimal=22)

castes = [
    ('CST_SC', 'SC', 'अनुसूचित जाति (SC)', 'अनुसूचित जाति (SC)'),
    ('CST_ST', 'ST', 'अनुसूचित जनजाति (ST)', 'अनुसूचित जनजाति (ST)'),
    ('CST_OBC', 'OBC', 'अन्य पिछड़ा वर्ग (OBC)', 'अन्य पिछड़ा वर्ग (OBC)'),
    ('CST_GEN', 'General', 'सामान्य (General)', 'सामान्य (General)')
]
cst_ids = [c[0] for c in castes]
for cid, en, hi, raj in castes:
    add_var(cid, 'AppVariables', 'Caste', 'Option', 'Enum', en, '', 'Caste Option', en, title_hi=hi, title_raj=raj)
add_var('Q_B_03_00', 'Survey', 'SocialCategory', 'QuestionPrompt, SectionB, Order:23', 'Enum', 'What is the social category?', 'Section B Q3', 'Social group', var_list=' , '.join(cst_ids), title_hi='सामाजिक वर्ग / जाति क्या है?', title_raj='जाति वर्ग कांई छै?', decimal=23)

educations = [
    ('EDU_ILLITERATE', 'Illiterate', 'निरक्षर', 'अणपढ़'),
    ('EDU_ILLITERATE_CALC', 'Illiterate but able to calculate', 'निरक्षर लेकिन हिसाब-किताब में सक्षम', 'अणपढ़ पण हिसाब-किताब जांणै'),
    ('EDU_5TH', 'Upto 5th', '5वीं तक', '5वीं तांई'),
    ('EDU_8TH', 'Upto 8th', '8वीं तक', '8वीं तांई'),
    ('EDU_10TH', 'Upto 10th', '10वीं तक', '10वीं तांई'),
    ('EDU_12TH', 'Upto 12th', '12वीं तक', '12वीं तांई'),
    ('EDU_DIPLOMA', 'Diploma', 'डिप्लोमा', 'डिप्लोमा'),
    ('EDU_GRADUATE', 'Graduate', 'स्नातक (Graduate)', 'कॉलेज पास (ग्रेजुएट)'),
    ('EDU_BED', 'B.Ed', 'बी.एड (B.Ed)', 'बी.एड')
]
edu_ids = [e[0] for e in educations]
for eid, en, hi, raj in educations:
    add_var(eid, 'AppVariables', 'Education', 'Option', 'Enum', en, '', 'Education Option', en, title_hi=hi, title_raj=raj)
add_var('Q_B_04_00', 'Survey', 'EducationStatus', 'QuestionPrompt, SectionB, Order:24', 'Enum', 'What is the education status?', 'Section B Q4', 'Education level', var_list=' , '.join(edu_ids), title_hi='शैक्षणिक योग्यता क्या है?', title_raj='कित्ती पढ़ाई-लिखाई कर राखी छै?', decimal=24)

add_var('Q_B_05_00', 'Survey', 'FamilyMemberCount', 'QuestionPrompt, SectionB, Order:25', 'Number', 'How many members are in the family?', 'Section B Q5', 'Total family count', title_hi='परिवार में कुल कितने सदस्य हैं?', title_raj='कुटुंब में कुल कित्ता जणा छै?', decimal=25)
add_var('Q_B_06_01', 'Survey', 'FamilyAdultsCount', 'QuestionPrompt, SectionB, Order:26', 'Number', 'Adults (Above 18)', 'Section B Q6.1', 'Adult count', title_hi='वयस्क (18 वर्ष से अधिक)', title_raj='बड़ा जणा (18 साल सूं बत्ता)', decimal=26)
add_var('Q_B_06_02', 'Survey', 'FamilyChildrenCount', 'QuestionPrompt, SectionB, Order:27', 'Number', 'Children', 'Section B Q6.2', 'Child count', title_hi='बच्चे (18 वर्ष से कम)', title_raj='टाबर (18 साल सूं कम)', decimal=27)
add_var('Q_B_06_03', 'Survey', 'FamilyTotalEarning', 'QuestionPrompt, SectionB, Order:28', 'Number', 'Total earning members', 'Section B Q6.3', 'Total earning count', title_hi='कुल कमाने वाले सदस्य', title_raj='कुल कमावणिया जणा', decimal=28)
add_var('Q_B_06_04', 'Survey', 'FamilyMaleEarning', 'QuestionPrompt, SectionB, Order:29', 'Number', 'Male earning members', 'Section B Q6.4', 'Male earning count', title_hi='पुरुष कमाने वाले सदस्य', title_raj='कमावणिया आदमी', decimal=29)
add_var('Q_B_06_05', 'Survey', 'FamilyFemaleEarning', 'QuestionPrompt, SectionB, Order:30', 'Number', 'Female earning members', 'Section B Q6.5', 'Female earning count', title_hi='महिला कमाने वाली सदस्य', title_raj='कमावणिया लुगायां', decimal=30)
add_var('Q_B_06_06', 'Survey', 'FamilyDisabledCount', 'QuestionPrompt, SectionB, Order:31', 'Number', 'Members with disability', 'Section B Q6.6', 'Disabled count', title_hi='दिव्यांग सदस्य', title_raj='दिव्यांग सदस्य', decimal=31)

income_sources = [
    ('INC_AGRI', 'Agricultural income', 'कृषि आय', 'खेती-बाड़ी री कमाई'),
    ('INC_SALARY', 'Fixed Salary', 'नियमित वेतन / नौकरी', 'पक्की तनख्वाह / नौकरी'),
    ('INC_WAGES', 'Wages', 'दैनिक मजदूरी', 'मजूरी'),
    ('INC_SELF_EMP', 'Self employed', 'स्वरोजगार', 'खुद रो रोजगार'),
    ('INC_NTFP', 'NTFP sale', 'लघु वनोपज (NTFP) बिक्री', 'जंगल री उपज बेचान'),
    ('INC_DAIRY', 'Dairying', 'डेयरी व्यवसाय', 'दूध-डेरी रो काम'),
    ('INC_ANIMAL_PROD', 'Animal products', 'पशु उत्पाद बिक्री', 'पशु उत्पाद बेचान'),
    ('INC_FAMILY_ENT', "Family/husband's enterprise", 'परिवार / पति का उद्यम', 'घर रो / धणी रो धंधो'),
    ('INC_RESP_ENT', "Respondent's enterprise", 'उत्तरदाता का स्वयं का उद्यम', 'थारो खुद रो धंधो'),
    ('INC_MNREGA', 'MNREGA', 'मनरेगा', 'नरेगा मजूरी'),
    ('INC_PENSION', 'Pension', 'पेंशन', 'पेंशन'),
    ('INC_RENT', 'Rent from properties', 'मकान / दुकान किराया', 'किरायो')
]
inc_src_ids = [i[0] for i in income_sources]
for iid, en, hi, raj in income_sources:
    add_var(iid, 'AppVariables', 'IncomeSource', 'Option', 'Enum', en, '', 'Income Source Option', en, title_hi=hi, title_raj=raj)
add_var('Q_B_07_00', 'Survey', 'FamilyIncomeSources', 'QuestionPrompt, SectionB, Order:32', 'EnumList', 'What are your family’s sources of income? (Multiselect)', 'Section B Q7', 'Income sources', var_list=' , '.join(inc_src_ids), title_hi='आपके परिवार की आय के क्या स्रोत हैं?', title_raj='थारे कुटुंब री कमाई रा कांई साधन छै?', decimal=32)

income_brackets = [
    ('INC_LT_80K', 'Less than Rs 80,000', 'रु 80,000 से कम', '80 हजार सूं कम'),
    ('INC_80K_120K', 'Rs 80,000 to Rs 1,20,000', 'रु 80,000 से रु 1,20,000', '80 हजार सूं 1 लाख 20 हजार'),
    ('INC_120K_160K', 'Rs 1,20,000 to Rs 1,60,000', 'रु 1,20,000 से रु 1,60,000', '1.20 लाख सूं 1.60 लाख'),
    ('INC_160K_200K', 'Rs 1,60,000 to Rs 2,00,000', 'रु 1,60,000 से रु 2,00,000', '1.60 लाख सूं 2 लाख'),
    ('INC_200K_240K', 'Rs 2,00,000 to Rs 2,40,000', 'रु 2,00,000 से रु 2,40,000', '2 लाख सूं 2.40 लाख'),
    ('INC_240K_280K', 'Rs 2,40,000 to Rs 2,80,000', 'रु 2,40,000 से रु 2,80,000', '2.40 लाख सूं 2.80 लाख'),
    ('INC_280K_320K', 'Rs 2,80,000 to Rs 3,20,000', 'रु 2,80,000 से रु 3,20,000', '2.80 लाख सूं 3.20 लाख'),
    ('INC_320K_360K', 'Rs 3,20,000 to Rs 3,60,000', 'रु 3,20,000 से रु 3,60,000', '3.20 लाख सूं 3.60 लाख'),
    ('INC_GT_360K', 'Above Rs 3,60,000', 'रु 3,60,000 से अधिक', '3 लाख 60 हजार सूं बत्ता')
]
inc_brk_ids = [b[0] for b in income_brackets]
for bid, en, hi, raj in income_brackets:
    add_var(bid, 'AppVariables', 'AnnualIncomeBracket', 'Option', 'Enum', en, '', 'Bracket Option', en, title_hi=hi, title_raj=raj)
add_var('Q_B_08_00', 'Survey', 'AnnualHouseholdIncome', 'QuestionPrompt, SectionB, Order:33', 'Enum', 'What is your annual household income and monetary benefits from all sources?', 'Section B Q8', 'Household income bracket', var_list=' , '.join(inc_brk_ids), title_hi='सभी स्रोतों से आपके परिवार की कुल वार्षिक आय कितनी है?', title_raj='सगळा साधनां सूं कुटुंब री साल भर री कमाई कित्ती छै?', decimal=33)

# ==========================================
# 4. SECTION C: ENTERPRISE OPERATIONS
# ==========================================
reasons_starting = [
    ('RSN_SETBACK', 'My family faced a financial setback, and I needed to earn', 'परिवार को आर्थिक संकट का सामना करना पड़ा और मुझे कमाना पड़ा', 'घर में तंगी आग्यी ही, इस वास्ते कमावणो पड़्यो'),
    ('RSN_RISING_EXP', 'Our expenses were rising, and my family needed an alternate source of income', 'खर्च बढ़ रहे थे और परिवार को अतिरिक्त आय की आवश्यकता थी', 'खर्चा बढ़ग्या हा और दूजी कमाई री जरूरत ही'),
    ('RSN_OWN_VENTURE', 'I always wanted to own/run my own business', 'मेरी हमेशा से अपना व्यवसाय चलाने की इच्छा थी', 'म्हारी खुद रो धंधो चलावण री इच्छा ही'),
    ('RSN_LEARNT_SKILL', 'I learnt the skill and wanted to start my own venture', 'मैंने हुनर सीखा और अपना काम शुरू करना चाहा', 'काम सीख्यो और खुद रो धंधो सुरू करयो'),
    ('RSN_FROM_WAGE', 'I was doing the same work as wage labour and later decided to start own venture', 'मैं मजदूरी करती थी और बाद में खुद का उद्यम शुरू करने का फैसला किया', 'पैली मजूरी करती ही, फेर खुद रो काम सुरू करयो'),
    ('RSN_ALL_SHG_LOAN', 'All SHG members were getting loans for enterprise so I also decided to take and start', 'सभी SHG सदस्यों को लोन मिल रहा था तो मैंने भी उद्यम शुरू किया', 'सगळी बाईयां लोन लेवती ही तो म्हैं भी लोन लियो'),
    ('RSN_CRP_ENCOURAGED', 'The OSF/SVEP CRP encouraged me to start the enterprise', 'OSF/SVEP सीआरपी ने मुझे उद्यम शुरू करने के लिए प्रेरित किया', 'सीआरपी दीदी हौसलो दियो और काम सुरू करवायो'),
    ('RSN_CLF_ENCOURAGED', 'The CLF encouraged me to start the enterprise', 'सीएलएफ ने मुझे उद्यम शुरू करने के लिए प्रेरित किया', 'सीएलएफ सूं प्रेरणा मिली')
]
rsn_ids = [r[0] for r in reasons_starting]
for rid, en, hi, raj in reasons_starting:
    add_var(rid, 'AppVariables', 'ReasonStarting', 'Option', 'Enum', en, '', 'Reason Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_01_00', 'Survey', 'ReasonsStartingBusiness', 'QuestionPrompt, SectionC, Order:34', 'EnumList', 'Reasons for starting the business? (Multiselect)', 'Section C Q1', 'Reasons for starting', var_list=' , '.join(rsn_ids), title_hi='व्यवसाय शुरू करने के क्या कारण रहे?', title_raj='काम-धंधो सुरू करण रा कांई कारण हा?', decimal=34)

biz_cycles = [
    ('CYC_REGULAR_HOURS', 'Operational for regular hours throughout the year', 'पूरे वर्ष नियमित घंटों तक संचालित', 'पूरे साल रोज बंध्या टेम पर खुलै'),
    ('CYC_WHENEVER_CUSTOMER', 'Operational whenever the customer arrives throughout the year', 'पूरे वर्ष जब भी ग्राहक आए तब संचालित', 'पूरे साल जद भी ग्राहक आवै जद खुलै'),
    ('CYC_BOTH_ROUND_YEAR', 'Both production and sales operational throughout the year', 'पूरे वर्ष उत्पादन और बिक्री दोनों चालू रहते हैं', 'पूरे साल माल बणावण और बेचण रो काम चालै'),
    ('CYC_ON_ORDER_ONLY', 'Production and sale only on receiving order', 'केवल ऑर्डर मिलने पर उत्पादन और बिक्री', 'ऑर्डर मिल्या पर ही माल बणावै और बेचै'),
    ('CYC_SEASONAL_ROUND', 'Seasonal production and sale throughout the year', 'मौसमी उत्पादन लेकिन पूरे साल बिक्री', 'मौसम में माल बणै पण पूरे साल बिकै'),
    ('CYC_FEW_MONTHS', 'Production and sale is limited to few months', 'उत्पादन और बिक्री कुछ महीनों तक सीमित', 'साल में बस कुछेक महिना ही काम चालै'),
    ('CYC_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजो कोई प्रकार')
]
cyc_ids = [c[0] for c in biz_cycles]
for cid, en, hi, raj in biz_cycles:
    add_var(cid, 'AppVariables', 'BusinessCycle', 'Option', 'Enum', en, '', 'Cycle Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_02_00', 'Survey', 'BusinessCycle', 'QuestionPrompt, SectionC, Order:35', 'Enum', 'Describe your business cycle?', 'Section C Q2', 'Cycle type', var_list=' , '.join(cyc_ids), title_hi='अपने व्यवसाय चक्र का विवरण दें', title_raj='थारो धंधो साल में किस तरियां चालै छै?', decimal=35)
add_var('Q_C_02_01', 'Survey', 'BusinessCycleOther', 'QuestionPrompt, SectionC, Order:36', 'Text', 'Specify other business cycle', 'Section C Q2.1', 'Other cycle description', title_hi='अन्य व्यवसाय चक्र का विवरण दें', title_raj='दूजे प्रकार रो ब्योरो लिखो', decimal=36)

places = [
    ('PLC_OWN', 'Own', 'स्वयं की जगह / दुकान', 'खुद री जगह'),
    ('PLC_RENTED', 'Rented', 'किराए की जगह / दुकान', 'किराये री जगह')
]
plc_ids = [p[0] for p in places]
for pid, en, hi, raj in places:
    add_var(pid, 'AppVariables', 'BusinessPlace', 'Option', 'Enum', en, '', 'Place Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_03_00', 'Survey', 'BusinessPlaceType', 'QuestionPrompt, SectionC, Order:37', 'Enum', 'What is the type of business place?', 'Section C Q3', 'Place type', var_list=' , '.join(plc_ids), title_hi='व्यवसाय स्थल का प्रकार क्या है?', title_raj='दुकान/काम री जगह खुद री छै या किराये री?', decimal=37)
add_var('Q_C_04_00', 'Survey', 'AnnualRent', 'QuestionPrompt, SectionC, Order:38', 'Decimal', 'If rented, what is annual rent? (Rs)', 'Section C Q4', 'Annual rent amount', title_hi='यदि किराए पर है, तो वार्षिक किराया कितना है? (रु)', title_raj='अगर किराये री छै, तो साल रो कित्तो भाड़ो लागै?', decimal=38)

locations = [
    ('LOC_CONVENIENT', 'Yes, my location is very convenient to attract customers', 'हाँ, ग्राहकों को आकर्षित करने के लिए स्थान बहुत सुविधाजनक है', 'हाँ, गिराहक आराम सूं आवै, बढ़िया जगह छै'),
    ('LOC_CHANGED_LOC', 'Yes, I changed my location to get the clients', 'हाँ, ग्राहकों के लिए मैंने अपना स्थान बदला', 'हाँ, गिराहकां खातर जगह बदली'),
    ('LOC_HOME_CANT_MOVE', 'No, but I operate from home and can’t move to other location', 'नहीं, लेकिन मैं घर से काम करती हूँ और दूसरी जगह नहीं जा सकती', 'ना, पण घर सूं काम करां, बाहर नीं जा सका'),
    ('LOC_AFFORD_ONLY', 'No, but I can afford only this space', 'नहीं, लेकिन मैं केवल इसी जगह का खर्च उठा सकती हूँ', 'ना, पण म्हारे बजट में आ ही जगह ही'),
    ('LOC_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजी कोई बात')
]
loc_ids = [l[0] for l in locations]
for lid, en, hi, raj in locations:
    add_var(lid, 'AppVariables', 'LocationConvenience', 'Option', 'Enum', en, '', 'Location Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_05_00', 'Survey', 'LocationConvenience', 'QuestionPrompt, SectionC, Order:39', 'Enum', 'Is the location of your premise convenient for your customers?', 'Section C Q5', 'Location convenience', var_list=' , '.join(loc_ids), title_hi='क्या आपके परिसर का स्थान ग्राहकों के लिए सुविधाजनक है?', title_raj='कांई थारी दुकान/जगह गिराहकां सारु सुभीते री छै?', decimal=39)
add_var('Q_C_05_01', 'Survey', 'LocationConvenienceOther', 'QuestionPrompt, SectionC, Order:40', 'Text', 'Specify other location convenience remark', 'Section C Q5.1', 'Other location remark', title_hi='स्थान सुविधा पर अन्य टिप्पणी', title_raj='जगह सारु दूजी बात लिखो', decimal=40)

# Labor Involvement Options
involvements = [
    ('INV_REGULAR', 'Regular', 'नियमित', 'रोज/नियमित'),
    ('INV_OCCASIONAL', 'Occasional', 'कभी-कभार', 'कदे-कदाई'),
    ('INV_ONLY_RESP', 'Only respondent', 'केवल उत्तरदाता', 'सिर्फ म्हैं खुद'),
    ('INV_NOT_RELEVANT', 'Not relevant', 'लागू नहीं', 'लागू कोनी')
]
inv_ids = [i[0] for i in involvements]
for iid, en, hi, raj in involvements:
    add_var(iid, 'AppVariables', 'InvolvementLevel', 'Option', 'Enum', en, '', 'Involvement Option', en, title_hi=hi, title_raj=raj)

# Labor Matrix Prompts
labor_activities = [
    ('Purchase', 'Purchase of material', 'सामग्री की खरीद', 'सामान खरीदण'),
    ('Prod', 'Production', 'उत्पादन / निर्माण', 'माल बणावण'),
    ('Serv', 'Servicing', 'सेवा / सर्विसिंग', 'सेवा रो काम'),
    ('Mktg', 'Social media marketing', 'सोशल मीडिया मार्केटिंग', 'मोबाइल/प्रचार काम'),
    ('Sale', 'Sale (shop/door to door/haat)', 'बिक्री (दुकान/हाट/मेला)', 'माल बेचण'),
    ('Record', 'Record keeping', 'हिसाब-किताब / रिकॉर्ड रखना', 'हिसाब-किताब रखण')
]
ord_idx = 41
for lkey, len_title, lhi, lraj in labor_activities:
    add_var(f'Q_C_06_{lkey}_INV', 'Survey', f'Labor_{lkey}_Involvement', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', f'{len_title} — Involvement of family members', f'Section C Q6 {lkey}', 'Involvement level', var_list=' , '.join(inv_ids), title_hi=f'{lhi} — परिवार के सदस्यों की भागीदारी', title_raj=f'{lraj} — घर रा जणां री भागीदारी', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_06_{lkey}_FAM', 'Survey', f'Labor_{lkey}_FamilyCount', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Number', f'{len_title} — Family members involved (#)', f'Section C Q6 {lkey}', 'Family count', title_hi=f'{lhi} — शामिल परिवार के सदस्य (संख्या)', title_raj=f'{lraj} — घर रा जणां री गिनती', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_06_{lkey}_HIRED', 'Survey', f'Labor_{lkey}_HiredCount', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Number', f'{len_title} — Hired help (#)', f'Section C Q6 {lkey}', 'Hired count', title_hi=f'{lhi} — रखे गए मजदूर/सहायक (संख्या)', title_raj=f'{lraj} — मजूरा री गिनती', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_06_{lkey}_AMT', 'Survey', f'Labor_{lkey}_AmountPaid', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{len_title} — Amount paid in last one year (Rs)', f'Section C Q6 {lkey}', 'Amount paid', title_hi=f'{lhi} — पिछले एक वर्ष में भुगतान की गई राशि (रु)', title_raj=f'{lraj} — साल भर में कित्ता पीसा दिया (रु)', decimal=ord_idx)
    ord_idx += 1

# Salary Bill
salary_bills = [
    ('SAL_NOT_REL', 'Not relevant', 'लागू नहीं', 'लागू कोनी'),
    ('SAL_UPTO_5K', 'Upto Rs 5000', 'रु 5,000 तक', '5 हजार तांई'),
    ('SAL_6K_10K', 'Rs 6000 to Rs 10,000', 'रु 6,000 से रु 10,000', '6 हजार सूं 10 हजार'),
    ('SAL_11K_30K', 'Rs 11,000 to Rs 30,000', 'रु 11,000 से रु 30,000', '11 हजार सूं 30 हजार'),
    ('SAL_31K_50K', 'Rs 31,000 to Rs 50,000', 'रु 31,000 से रु 50,000', '31 हजार सूं 50 हजार'),
    ('SAL_51K_70K', 'Rs 51,000 to Rs 70,000', 'रु 51,000 से रु 70,000', '51 हजार सूं 70 हजार'),
    ('SAL_71K_90K', 'Rs 71,000 to Rs 90,000', 'रु 71,000 से रु 90,000', '71 हजार सूं 90 हजार'),
    ('SAL_90K_110K', 'Rs 90,000 to Rs 1,10,000', 'रु 90,000 से रु 1,10,000', '90 हजार सूं 1.10 लाख'),
    ('SAL_111K_130K', 'Rs 1,11,000 to Rs 1,30,000', 'रु 1,11,000 से रु 1,30,000', '1.11 लाख सूं 1.30 लाख'),
    ('SAL_131K_150K', 'Rs 1,31,000 to Rs 1,50,000', 'रु 1,31,000 से रु 1,50,000', '1.31 लाख सूं 1.50 लाख'),
    ('SAL_GT_150K', 'Above Rs 1,50,000', 'रु 1,50,000 से अधिक', '1.50 लाख सूं बत्ता')
]
sal_ids = [s[0] for s in salary_bills]
for sid, en, hi, raj in salary_bills:
    add_var(sid, 'AppVariables', 'SalaryBill', 'Option', 'Enum', en, '', 'Salary Bill Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_07_00', 'Survey', 'AnnualSalaryBill', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'What is your annual salary bill (paid to hired help)?', 'Section C Q7', 'Salary bill bracket', var_list=' , '.join(sal_ids), title_hi='मजदूरों/सहायकों को दिया जाने वाला वार्षिक वेतन कितना है?', title_raj='मजूरां नें साल भर में कुल कित्ती पगार दी?', decimal=ord_idx)
ord_idx += 1

# Percent brackets (0, 25, 50, 75, 100)
pct_opts = [
    ('PCT_0', '0%', '0%', '0%'),
    ('PCT_25', '25%', '25%', '25%'),
    ('PCT_50', '50%', '50%', '50%'),
    ('PCT_75', '75%', '75%', '75%'),
    ('PCT_100', '100%', '100%', '100%')
]
pct_ids = [p[0] for p in pct_opts]
for pid, en, hi, raj in pct_opts:
    add_var(pid, 'AppVariables', 'PercentageLevel', 'Option', 'Enum', en, '', 'Percentage Option', en, title_hi=hi, title_raj=raj)

sourcing_places = [
    ('NearbyTown', 'Nearby town/district', 'पास का कस्बा / जिला', 'नेड़लो शहर / जिलो'),
    ('Jaipur', 'Jaipur', 'जयपुर', 'जयपुर'),
    ('OutsideState', 'Outside the state', 'राज्य के बाहर से', 'राज सूं बाहर'),
    ('Online', 'Order online (Amazon/Meesho)', 'ऑनलाइन ऑर्डर (अमेज़न/मीशो)', 'ऑनलाइन मंगवायो'),
    ('WhatsApp', 'Order using WhatsApp from existing sources', 'व्हाट्सएप से पुराने सप्लायर से ऑर्डर', 'व्हाट्सएप सूं मंगवायो')
]
for skey, sen, shi, sraj in sourcing_places:
    add_var(f'Q_C_08_{skey}', 'Survey', f'Sourcing_{skey}_Pct', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', f'Material sourcing % from {sen}', f'Section C Q8 {skey}', 'Sourcing %', var_list=' , '.join(pct_ids), title_hi=f'{shi} से सामग्री खरीद का %', title_raj=f'{sraj} सूं माल मंगावण रो %', decimal=ord_idx)
    ord_idx += 1

# Marketing Methods
marketing_opts = [
    ('MKT_SHOP_ONLY', 'My shop is the only place where I talk about my products/services', 'दुकान ही एकमात्र जगह है जहाँ मैं प्रचार करती हूँ', 'दुकान पर ही माल री बात करां'),
    ('MKT_NAME_BOARD', 'I have name board outside my premises with details', 'दुकान के बाहर नाम बोर्ड लगा रखा है', 'दुकान बारै नाम रो बोर्ड लगा राख्यो छै'),
    ('MKT_DOOR_TO_DOOR', 'I visit local traders/shopkeepers and do door to door selling', 'व्यापारियों के पास जाती हूँ और घर-घर जाकर बेचती हूँ', 'घर-घर जा र माल बेचां'),
    ('MKT_SHG_MEETINGS', 'I talk about my products/services in SHG meetings', 'SHG बैठकों में अपने उत्पादों/सेवाओं की बात करती हूँ', 'समूह री मीटिंग में बात करां'),
    ('MKT_TRADERS_SAMPLES', 'I visit local traders/shopkeepers with samples of my products', 'व्यापारियों को उत्पाद के नमूने (सैंपल) दिखाती हूँ', 'व्यापारियां नें सैंपल दिखावां'),
    ('MKT_WAIT_ENQUIRIES', 'I wait for people to make enquiries', 'ग्राहकों के स्वयं पूछताछ करने का इंतजार करती हूँ', 'गिराहक आवै जद ही बात करां'),
    ('MKT_DONT_KNOW_HOW', 'I do not know how to market my products/services', 'मुझे नहीं पता कि मार्केटिंग/प्रचार कैसे करें', 'म्हानें प्रचार करणो कोनी आवै'),
    ('MKT_NO_NEED', "I don't feel the need to market my products/services", 'मुझे प्रचार करने की आवश्यकता महसूस नहीं होती', 'प्रचार करण री जरूरत कोनी लागै'),
    ('MKT_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजो कोई तरीको')
]
mkt_ids = [m[0] for m in marketing_opts]
for mid, en, hi, raj in marketing_opts:
    add_var(mid, 'AppVariables', 'MarketingMethod', 'Option', 'Enum', en, '', 'Marketing Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_09_00', 'Survey', 'MarketingMethods', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'EnumList', 'How do you market your products/services? (Multiselect)', 'Section C Q9', 'Marketing methods', var_list=' , '.join(mkt_ids), title_hi='आप अपने उत्पादों/सेवाओं का प्रचार कैसे करती हैं?', title_raj='माल बेचण खातर कांई तरीको अपनावो छो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_09_01', 'Survey', 'MarketingMethodsOther', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Text', 'Specify other marketing method', 'Section C Q9.1', 'Other marketing remark', title_hi='अन्य मार्केटिंग तरीके का विवरण दें', title_raj='दूजे तरीके रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1

# Seasonal Sales
seasonal_opts = [
    ('SEA_NOT_REL', 'Not relevant', 'लागू नहीं', 'लागू कोनी'),
    ('SEA_LAST_YEAR_ORDERS', 'I produce as per last year sales and wait for orders', 'पिछले साल की बिक्री के अनुसार बनाकर ऑर्डर का इंतजार करती हूँ', 'पिछली साल जेड़ो बणा र ऑर्डर री बाट देखां'),
    ('SEA_DOOR_TO_DOOR', 'I visit local traders/shopkeepers and do door to door selling', 'व्यापारियों से मिलती हूँ और घर-घर जाकर बेचती हूँ', 'घर-घर जा र बेचां'),
    ('SEA_PRIOR_WEEKS_ORDERS', 'Visit local traders few weeks prior to peak season and get orders', 'सीजन से कुछ हफ्ते पहले व्यापारियों से ऑर्डर लेती हूँ', 'सीजन सूं पैली आर्डर लेवां'),
    ('SEA_LOCAL_HAAT', 'I sell in local haat/weekly market', 'स्थानीय हाट/साप्ताहिक बाजार में बेचती हूँ', 'हाट/साप्ताहिक बाजार में बेचां'),
    ('SEA_SARAS_FAIR', 'I sell in Saras fair', 'सरस मेले में बेचती हूँ', 'सरस मेला में बेचां'),
    ('SEA_ONLINE_PLATFORM', 'I use online platforms to sell (Specify platform)', 'ऑनलाइन प्लेटफॉर्म से बेचती हूँ (प्लेटफॉर्म लिखें)', 'ऑनलाइन बेचूं'),
    ('SEA_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजो कोई तरीको')
]
sea_ids = [s[0] for s in seasonal_opts]
for sid, en, hi, raj in seasonal_opts:
    add_var(sid, 'AppVariables', 'SeasonalSalesMethod', 'Option', 'Enum', en, '', 'Seasonal Sales Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_10_00', 'Survey', 'SeasonalSalesMethod', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'In case of seasonal production, how do you sell your products/services?', 'Section C Q10', 'Seasonal sales method', var_list=' , '.join(sea_ids), title_hi='मौसमी उत्पादन की स्थिति में, आप अपने उत्पादों/सेवाओं को कैसे बेचती हैं?', title_raj='सीजन रा टेम माल किकण बेचो छो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_10_01', 'Survey', 'SeasonalSalesOnlinePlatform', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Text', 'Specify online platform used', 'Section C Q10.1', 'Platform name', title_hi='उपयोग किए गए ऑनलाइन प्लेटफॉर्म का नाम लिखें', title_raj='ऑनलाइन साधन रो नाम लिखो', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_10_02', 'Survey', 'SeasonalSalesOther', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Text', 'Specify other seasonal method', 'Section C Q10.2', 'Other seasonal remark', title_hi='अन्य मौसमी बिक्री तरीके का विवरण दें', title_raj='दूजे तरीके रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1

# Social Media For Marketing
social_mktg_opts = [
    ('SMM_WHATSAPP_ORDERS', 'I regularly share images on whatsapp to get orders', 'ऑर्डर पाने के लिए नियमित व्हाट्सएप पर फोटो भेजती हूँ', 'व्हाट्सएप पर फोटो भेज र आर्डर लेवां'),
    ('SMM_INSTA_ORDERS', 'I regularly share images/reels on instagram to get orders', 'ऑर्डर पाने के लिए इंस्टाग्राम पर फोटो/रील शेयर करती हूँ', 'इंस्टाग्राम पर रील/फोटो शेयर करां'),
    ('SMM_NO_SMARTPHONE', 'It is important but I don’t have access to smart phone', 'महत्वपूर्ण है लेकिन मेरे पास स्मार्टफोन नहीं है', 'जरूरी छै पण म्हारे कनै स्मार्टफोन कोनी'),
    ('SMM_DONT_KNOW_USE', 'I do not use because I don’t know how to use whatsapp/instagram', 'उपयोग नहीं करती क्योंकि व्हाट्सएप/इंस्टाग्राम चलाना नहीं आता', 'व्हाट्सएप/इंस्टा चलावणो कोनी आवै'),
    ('SMM_NO_TIME_LEARN', "I don't have time to learn and use social media", 'सीखने और उपयोग करने का समय नहीं है', 'सीखण रो टेम कोनी'),
    ('SMM_DONT_WANT', "I don't want to use social media", 'मैं सोशल मीडिया का उपयोग नहीं करना चाहती', 'सोशल मीडिया कोनी चलावणो'),
    ('SMM_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजी कोई बात')
]
smm_ids = [s[0] for s in social_mktg_opts]
for sid, en, hi, raj in social_mktg_opts:
    add_var(sid, 'AppVariables', 'SocialMediaForMarketing', 'Option', 'Enum', en, '', 'SMM Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_11_00', 'Survey', 'SocialMediaForMarketing', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'Do you use social media for marketing?', 'Section C Q11', 'Social media marketing', var_list=' , '.join(smm_ids), title_hi='क्या आप मार्केटिंग के लिए सोशल मीडिया का उपयोग करती हैं?', title_raj='कांई थे प्रचार खातर सोशल मीडिया चलावो छो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_11_01', 'Survey', 'SocialMediaForMarketingOther', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Text', 'Specify other social media reason', 'Section C Q11.1', 'Other SMM remark', title_hi='अन्य सोशल मीडिया कारण का विवरण दें', title_raj='दूजी बात लिखो', decimal=ord_idx)
ord_idx += 1

# Sales Channels % Breakdown
sales_channels = [
    ('Online', 'Online platforms', 'ऑनलाइन प्लेटफॉर्म', 'ऑनलाइन साधन'),
    ('WhatsApp', 'WhatsApp', 'व्हाट्सएप', 'व्हाट्सएप'),
    ('Instagram', 'Instagram', 'इंस्टाग्राम', 'इंस्टाग्राम'),
    ('Premise', 'Your premise / shop', 'आपकी दुकान / परिसर', 'खुद री दुकान'),
    ('Traders', 'Local traders / shopkeepers', 'स्थानीय व्यापारी / दुकानदार', 'गांम/शहर रा व्यापारी'),
    ('Haat', 'Local haat / weekly market', 'स्थानीय हाट / बाजार', 'हाट/साप्ताहिक बाजार'),
    ('Saras', 'Saras fair', 'सरस मेला', 'सरस मेला')
]
for ckey, cen, chi, craj in sales_channels:
    add_var(f'Q_C_12_{ckey}', 'Survey', f'SalesChannel_{ckey}_Pct', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', f'% products/services sold through {cen}', f'Section C Q12 {ckey}', 'Sales channel %', var_list=' , '.join(pct_ids), title_hi=f'{chi} के माध्यम से बिक्री का %', title_raj=f'{craj} सूं बिक्री रो %', decimal=ord_idx)
    ord_idx += 1

# Record keeping habits
record_habits = [
    ('RKH_ALWAYS_DONE', 'Yes, I have always been doing it', 'हाँ, मैं हमेशा से लिखित हिसाब रखती आई हूँ', 'हाँ, म्हैं पैली सूं ही हिसाब राखूं'),
    ('RKH_AFTER_CRP_TRAIN', 'Yes, I started doing after being trained by OSF/SVEP CRP', 'हाँ, OSF/SVEP सीआरपी से प्रशिक्षण मिलने के बाद शुरू किया', 'हाँ, सीआरपी दीदी सिखायो जद सूं हिसाब राखूं'),
    ('RKH_FAMILY_MAINTAINS', 'Yes, my family member maintains but I don’t check', 'हाँ, परिवार का सदस्य हिसाब रखता है पर मैं नहीं देखती', 'घर रा जणां हिसाब राखै, म्हैं नीं देखूं'),
    ('RKH_HIRED_HELP', 'Yes, I have hired help to do that', 'हाँ, हिसाब रखने के लिए मुनीम/सहायक रखा है', 'हिसाब सारु सहायक राख्यो छै'),
    ('RKH_NOT_REGULAR', 'I don’t record regularly', 'नियमित रूप से हिसाब नहीं रखती', 'रोज हिसाब कोनी राखूं'),
    ('RKH_NO_RECORD', 'I don’t maintain any records at all', 'मैं कोई लिखित रिकॉर्ड नहीं रखती', 'कोई हिसाब कोनी राखूं')
]
rkh_ids = [r[0] for r in record_habits]
for rid, en, hi, raj in record_habits:
    add_var(rid, 'AppVariables', 'RecordKeepingHabit', 'Option', 'Enum', en, '', 'Habit Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_13_00', 'Survey', 'RecordKeepingHabit', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'Do you maintain written records of business transactions?', 'Section C Q13', 'Record keeping habit', var_list=' , '.join(rkh_ids), title_hi='क्या आप व्यावसायिक लेन-देन का लिखित रिकॉर्ड रखती हैं?', title_raj='कांई थे दुकान रो लिख्योड़ो हिसाब राखो छो?', decimal=ord_idx)
ord_idx += 1

record_tools = [
    ('RKT_RECEIPT_BILLS', 'Receipt book / bills', 'रसीद बुक / बिल', 'बिल बही / रसीद'),
    ('RKT_DAILY_DIARY', 'Maintain daily diary', 'दैनिक डायरी / बही-खाता', 'रोज री डायरी / खाता'),
    ('RKT_CRP_DIARY', 'Maintain daily diary as taught by OSF/SVEP CRP', 'OSF/SVEP सीआरपी द्वारा सिखाई गई डायरी', 'सीआरपी दीदी री बताई डायरी'),
    ('RKT_DIGITAL_APPS', 'Maintain digital records using Mera Bill, Bahi Khata app', 'मेरा बिल, बही खाता जैसे डिजिटल ऐप में दर्ज करती हूँ', 'मोबाइल ऐप (मेरा बिल, बही खाता) में'),
    ('RKT_NOT_REGULAR', 'Don’t record regularly', 'नियमित रिकॉर्ड नहीं रखती', 'नियमित कोनी राखूं'),
    ('RKT_FAMILY_BOOK', 'My family member maintains a book', 'परिवार का सदस्य एक बही में लिखता है', 'घर रा जणां पोथी में लिखै'),
    ('RKT_NO_RECORD', 'I don’t maintain any record', 'कोई रिकॉर्ड नहीं रखती', 'कोई हिसाब कोनी राखूं'),
    ('RKT_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजो कोई साधन')
]
rkt_ids = [t[0] for t in record_tools]
for tid, en, hi, raj in record_tools:
    add_var(tid, 'AppVariables', 'RecordKeepingMethod', 'Option', 'Enum', en, '', 'Method Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_14_00', 'Survey', 'RecordKeepingMethod', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'How do you maintain business transactions?', 'Section C Q14', 'Transaction tool', var_list=' , '.join(rkt_ids), title_hi='आप व्यावसायिक लेन-देन का हिसाब कैसे रखती हैं?', title_raj='लेन-देन रो हिसाब किकण राखो छो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_14_01', 'Survey', 'RecordKeepingOther', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Text', 'Specify other record keeping method', 'Section C Q14.1', 'Other tool remark', title_hi='अन्य हिसाब पद्धति का विवरण दें', title_raj='दूजी पद्धति रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1

# Turnover & Income Matrix (Peak, Average, Lean)
seasons = [
    ('Peak', 'Peak season', 'पीक सीजन (सबसे अधिक बिक्री)', 'घणी बिक्री रो सीजन'),
    ('Avg', 'Average season', 'औसत सीजन (सामान्य बिक्री)', 'साधारण सीजन'),
    ('Lean', 'Lean season', 'लीन सीजन (मंदी का समय)', 'मंदी रो टेम')
]
for skey, sen, shi, sraj in seasons:
    add_var(f'Q_C_15_{skey}_MTH', 'Survey', f'Turnover_{skey}_Months', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Number', f'{sen} — Duration in months (count)', f'Section C Q15 {skey}', 'Months count', title_hi=f'{shi} — अवधि (महीनों में)', title_raj=f'{sraj} — कित्ता महिना चालै?', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_15_{skey}_SALES', 'Survey', f'Turnover_{skey}_Sales', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{sen} — Monthly sales (Rs)', f'Section C Q15 {skey}', 'Monthly sales', title_hi=f'{shi} — औसत मासिक बिक्री (रु)', title_raj=f'{sraj} — महिना री बिक्री (रु)', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_15_{skey}_PROFIT', 'Survey', f'Turnover_{skey}_Profit', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{sen} — Monthly income / profit (Rs)', f'Section C Q15 {skey}', 'Monthly profit', title_hi=f'{shi} — मासिक शुद्ध आय / मुनाफा (रु)', title_raj=f'{sraj} — महिना रो नफो/मुनाफो (रु)', decimal=ord_idx)
    ord_idx += 1

add_var('Q_C_16_00', 'Survey', 'InitialStartCapital', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', 'With how much money did you start the enterprise? (Rs)', 'Section C Q16', 'Start capital', title_hi='आपने कितने रुपयों से उद्यम शुरू किया था? (रु)', title_raj='काम-धंधो सुरू करण खातर कित्ता पीसा लगाया हा? (रु)', decimal=ord_idx)
ord_idx += 1

start_sources = [
    ('CAP_MY_SAVINGS', 'My savings', 'मेरी अपनी बचत', 'म्हारी खुद री बचत'),
    ('CAP_HUSBAND', 'Husband funded', 'पति द्वारा वित्तपोषित', 'धणी पीसा दिया'),
    ('CAP_FAMILY_FUND', 'Other family members funded', 'परिवार के अन्य सदस्यों ने दिया', 'घर रा दूजा जणा दिया'),
    ('CAP_SHG_LOAN', 'SHG loan', 'SHG से ऋण लिया', 'समूह सूं लोन लियो'),
    ('CAP_SVEP_LOAN', 'SVEP loan', 'SVEP ऋण लिया', 'SVEP रो लोन लियो'),
    ('CAP_FAM_LOAN', 'Loan from family member', 'रिश्तेदारों से उधार', 'नातेदारां सूं उधार'),
    ('CAP_MONEYLENDER', 'Loan from moneylender', 'साहूकार से कर्ज', 'साहूकार सूं ब्याज पर'),
    ('CAP_NBFI', 'Loan from NBFI / MFI', 'NBFI / माइक्रोफाइनेंस ऋण', 'कंपनी रो लोन'),
    ('CAP_GOLD_MORTGAGE', 'Mortgaged gold/silver', 'सोना/चांदी गिरवी रखकर', 'सोना-चांदी गिरवी रख र')
]
cap_src_ids = [c[0] for c in start_sources]
for cid, en, hi, raj in start_sources:
    add_var(cid, 'AppVariables', 'InitialCapitalSource', 'Option', 'Enum', en, '', 'Source Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_17_00', 'Survey', 'InitialCapitalArranged', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'How did you arrange this amount?', 'Section C Q17', 'Arranging start capital', var_list=' , '.join(cap_src_ids), title_hi='आपने यह राशि कैसे जुटाई थी?', title_raj='आ प्रारंभिक पूंजी किकण भेळी करी ही?', decimal=ord_idx)
ord_idx += 1

shg_helps = [
    ('SHG_TRAINING', 'Attended the skill training offered by SHG', 'SHG द्वारा आयोजित कौशल प्रशिक्षण में भाग लिया', 'समूह रो हुनर/काम सीखण वाळो प्रशिक्षण लियो'),
    ('SHG_MEETINGS_SCOPE', 'Got information about the scope of business from SHG meetings', 'SHG बैठकों से व्यवसाय की संभावनाओं की जानकारी मिली', 'समूह री मीटिंग सूं धंधे री समझ मिली'),
    ('SHG_SUBSIDY', 'Got subsidy/grant due to SHG', 'SHG के माध्यम से सरकारी सब्सिडी / अनुदान मिला', 'समूह सूं सब्सिडी/अनुदान मिल्यो'),
    ('SHG_INITIATE_LOAN', 'Took loan from SHG to buy material to initiate business', 'शुरुआती सामान खरीदने हेतु SHG से ऋण लिया', 'काम सुरू करण खातर समूह सूं लोन लियो'),
    ('SHG_REGULAR_LOANS', 'Take loans from SHG regularly as per business requirements', 'व्यावसायिक आवश्यकतानुसार SHG से नियमित ऋण लेती हूँ', 'काम वास्ते समूह सूं नियमित लोन लेवां'),
    ('SHG_CRP_GUIDED', 'OSF/SVEP CRP guided me in setting up the business', 'OSF/SVEP सीआरपी ने व्यवसाय स्थापित करने में मार्गदर्शन दिया', 'सीआरपी दीदी काम जमावण में मदद करी'),
    ('SHG_MUDRA_LOAN', 'OSF/SVEP CRP helped me to get Mudra loan', 'OSF/SVEP सीआरपी ने मुद्रा लोन दिलाने में मदद की', 'सीआरपी दीदी मुद्रा लोन करवायो'),
    ('SHG_BANK_LOAN', 'OSF/SVEP CRP helped me to get bank loan', 'OSF/SVEP सीआरपी ने बैंक ऋण दिलाने में मदद की', 'सीआरपी दीदी बैंक सूं लोन करवायो')
]
shg_help_ids = [h[0] for h in shg_helps]
for hid, en, hi, raj in shg_helps:
    add_var(hid, 'AppVariables', 'SHGAssistance', 'Option', 'Enum', en, '', 'SHG Assistance Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_18_00', 'Survey', 'SHGAssociationAssistance', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'EnumList', 'How has the SHG association helped in your enterprise? (Multiselect)', 'Section C Q18', 'SHG benefits', var_list=' , '.join(shg_help_ids), title_hi='SHG से जुड़ाव ने आपके उद्यम में किस प्रकार सहायता की?', title_raj='समूह सूं जुड़बा सूं थारे काम-धंधे में कांई मदद मिली?', decimal=ord_idx)
ord_idx += 1

# Loan Usages (Q20)
loan_usages = [
    ('USE_SEED_CAPITAL', 'Seed capital to buy material and set up shop', 'दुकान स्थापित करने हेतु सामग्री व बीज पूंजी', 'दुकान सुरू करण और माल लेवण री पूंजी'),
    ('USE_NEW_MACHINE', 'Buy new machine to increase production (eg: sewing machine)', 'उत्पादन क्षमता बढ़ाने हेतु नई मशीन (उदा. सिलाई मशीन)', 'नई मशीन खरीदी (जैसें सिलाई मशीन)'),
    ('USE_ASSETS_STORE', 'Buy assets to store and sell new products (eg: fridge)', 'उत्पाद रखने व बेचने हेतु उपकरण (उदा. फ्रिज)', 'फ्रिज या दूजो साधन खरीद्यो'),
    ('USE_EXPAND_SPACE', 'Get additional space to expand business (eg: flour mill)', 'व्यवसाय विस्तार हेतु अतिरिक्त जगह (उदा. चक्की का कमरा)', 'दुकान/काम बढ़ावण सारु जगह ली'),
    ('USE_RANGE_VARIETY', 'Buy more material to increase product range', 'उत्पादों की विविधता बढ़ाने हेतु अधिक सामग्री', 'नवा-नवा माल री वैरायटी लाई'),
    ('USE_SCALE_VOLUME', 'Buy more material/inputs to increase scale', 'व्यापार का पैमाना और स्टॉक बढ़ाने हेतु', 'काम रो दायरो बढ़ायो'),
    ('USE_VEHICLE', 'Buy vehicle to access new market', 'माल लाने-ले जाने हेतु वाहन खरीदा', 'गाड़ी/वाहन खरीद्यो'),
    ('USE_TRANSPORT', 'Access better transport services', 'बेहतर परिवहन सेवाओं की व्यवस्था', 'भाड़ा/गाड़ी री व्यवस्था करी'),
    ('USE_SMARTPHONE', 'Buy mobile phone to promote/sell online', 'ऑनलाइन प्रचार व बिक्री हेतु मोबाइल फोन खरीदा', 'स्मार्टफोन खरीद्यो'),
    ('USE_OTHER', 'Any other, specify', 'अन्य कोई उपयोग', 'दूजो कोई उपयोग'),
    ('USE_NOT_USED', 'Not used the source', 'इस स्रोत का उपयोग नहीं किया', 'यो साधन काम कोनी लियो')
]
use_ids = [u[0] for u in loan_usages]
for uid, en, hi, raj in loan_usages:
    add_var(uid, 'AppVariables', 'LoanUsagePurpose', 'Option', 'Enum', en, '', 'Usage Option', en, title_hi=hi, title_raj=raj)

# Capital Arranged Over Time (14 Sources)
cap_sources_14 = [
    ('OwnSavings', 'Own Savings', 'अपनी बचत', 'खुद री बचत'),
    ('Family', 'Financed by family member', 'परिवार के सदस्य द्वारा वित्तपोषित', 'घर रा जणा दिया'),
    ('Profit', 'Profit from business', 'व्यवसाय का मुनाफा', 'धंधे रो मुनाफो'),
    ('MortgGold', 'Mortgaged gold/silver', 'सोना/चांदी गिरवी रखकर', 'सोना-चांदी गिरवी रख र'),
    ('SoldGold', 'Sold gold/silver', 'सोना/चांदी बेचकर', 'सोना-चांदी बेच र'),
    ('FamLoan', 'Loan from family', 'रिश्तेदारों से कर्ज', 'नातेदारां सूं लोन'),
    ('Moneylender', 'Loan from moneylender', 'साहूकार से ऋण', 'साहूकार सूं कर्ज'),
    ('SHGLoan', 'Loan from SHG', 'SHG से ऋण', 'समूह सूं लोन'),
    ('OSFSVEPLoan', 'Loan from OSF/SVEP', 'OSF/SVEP से ऋण', 'OSF/SVEP सूं लोन'),
    ('OSFSubsidy', 'Subsidy/grant under OSF/SVEP', 'OSF/SVEP के तहत सब्सिडी/अनुदान', 'योजना री सब्सिडी/अनुदान'),
    ('PrivSaving', 'Loan from private saving groups/BC', 'निजी बचत समूह / बीसी से कर्ज', 'निजी बीसी/ग्रुप सूं लोन'),
    ('NBFC', 'Loan from NBFC', 'NBFC / माइक्रोफाइनेंस ऋण', 'कंपनी रो लोन'),
    ('Mudra', 'Mudra loan', 'मुद्रा लोन', 'मुद्रा लोन'),
    ('BankLoan', 'Loan from banks', 'बैंक ऋण', 'बैंक सूं लोन')
]
for skey, sen, shi, sraj in cap_sources_14:
    add_var(f'Q_C_19_{skey}_YR1', 'Survey', f'Cap_{skey}_Yr1', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{sen} — First year amount (Rs)', f'Section C Q19 {skey}', 'Yr1 amount', title_hi=f'{shi} — पहले वर्ष की राशि (रु)', title_raj=f'{sraj} — पैले साल कित्ता पीसा (रु)', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_19_{skey}_MID', 'Survey', f'Cap_{skey}_Mid', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{sen} — In-between years amount (Rs)', f'Section C Q19 {skey}', 'Mid years amount', title_hi=f'{shi} — बीच के वर्षों की राशि (रु)', title_raj=f'{sraj} — बिचला सालां में कित्ता पीसा (रु)', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_19_{skey}_CUR', 'Survey', f'Cap_{skey}_Cur', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{sen} — Current year (2026-27) amount (Rs)', f'Section C Q19 {skey}', 'Current year amount', title_hi=f'{shi} — चालू वर्ष (2026-27) की राशि (रु)', title_raj=f'{sraj} — इबके साल (2026-27) में कित्ता पीसा (रु)', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_19_{skey}_PEN', 'Survey', f'Cap_{skey}_Pending', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{sen} — Amount pending / balance (Rs)', f'Section C Q19 {skey}', 'Pending amount', title_hi=f'{shi} — बकाया राशि (रु)', title_raj=f'{sraj} — बाकी कित्ता पीसा देणा छै (रु)', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_20_{skey}_USE', 'Survey', f'Cap_{skey}_Usage', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', f'{sen} — How loan used in business', f'Section C Q20 {skey}', 'Loan usage purpose', var_list=' , '.join(use_ids), title_hi=f'{shi} — ऋण का व्यवसाय में क्या उपयोग किया गया?', title_raj=f'{sraj} — लोन रा पीसा कांई काम में लगाया?', decimal=ord_idx)
    ord_idx += 1

income_increases = [
    ('INC_UPTO_2K', 'Upto Rs 2000', 'रु 2,000 तक', '2 हजार तांई'),
    ('INC_2K_3K', 'Rs 2000 to Rs 3000', 'रु 2,000 से रु 3,000', '2 हजार सूं 3 हजार'),
    ('INC_3K_4K', 'Rs 3000 to Rs 4000', 'रु 3,000 से रु 4,000', '3 हजार सूं 4 हजार'),
    ('INC_4K_5K', 'Rs 4000 to Rs 5000', 'रु 4,000 से रु 5,000', '4 हजार सूं 5 हजार'),
    ('INC_5K_6K', 'Rs 5000 to Rs 6000', 'रु 5,000 से रु 6,000', '5 हजार सूं 6 हजार'),
    ('INC_GT_6K', 'Above Rs 6000', 'रु 6,000 से अधिक', '6 हजार सूं बत्ता'),
    ('INC_CANT_SAY', 'Cant say', 'कह नहीं सकती', 'कह नीं सका / ठा कोनी')
]
inc_inc_ids = [i[0] for i in income_increases]
for iid, en, hi, raj in income_increases:
    add_var(iid, 'AppVariables', 'IncomeIncreaseBracket', 'Option', 'Enum', en, '', 'Increase Bracket Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_21_00', 'Survey', 'MonthlyIncomeIncreaseByOSFSVEP', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Enum', 'Amount by which monthly income increased directly due to OSF/SVEP loans', 'Section C Q21', 'Direct income rise', var_list=' , '.join(inc_inc_ids), title_hi='OSF/SVEP ऋण से हुए बदलावों के कारण औसत मासिक आय में कितनी वृद्धि हुई?', title_raj='लोन मिल्या पाछै महिना री कमाई कित्ती बढ़गी?', decimal=ord_idx)
ord_idx += 1

# Trajectory Changes (Q22)
trajectory_metrics = [
    ('Sales', 'Average sales/month', 'औसत मासिक बिक्री', 'महिना री बिक्री'),
    ('Income', 'Average monthly income', 'औसत मासिक आय', 'महिना री कमाई'),
    ('TradeStock', 'In case of trading, value of inventory/stock', 'ट्रेडिंग में स्टॉक/इन्वेंटरी का मूल्य', 'ट्रेडिंग माल रो मोल'),
    ('ProdInputs', 'In case of production, value of stock of inputs', 'उत्पादन में कच्चे माल (इनपुट) का मूल्य', 'कच्चा माल रो मोल'),
    ('ProdFinished', 'In case of production, value of finished products', 'उत्पादन में तैयार माल का मूल्य', 'तैयार माल रो मोल'),
    ('ServAssets', 'In case of servicing, value of enterprise related assets', 'सर्विसिंग में कार्य-संबंधित परिसंपत्तियों का मूल्य', 'सर्विसिंग औजार/मशीन रो मोल')
]
for tkey, ten, thi, traj in trajectory_metrics:
    add_var(f'Q_C_22_{tkey}_YR1', 'Survey', f'Trajectory_{tkey}_Yr1', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{ten} (First Year) [Rs]', f'Section C Q22 {tkey}', 'First year metric', title_hi=f'{thi} (पहले वर्ष में) [रु]', title_raj=f'{traj} (पैले साल में) [रु]', decimal=ord_idx)
    ord_idx += 1
    add_var(f'Q_C_22_{tkey}_CUR', 'Survey', f'Trajectory_{tkey}_Cur', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', f'{ten} (Current Year) [Rs]', f'Section C Q22 {tkey}', 'Current year metric', title_hi=f'{thi} (वर्तमान वर्ष में) [रु]', title_raj=f'{traj} (इबके साल में) [रु]', decimal=ord_idx)
    ord_idx += 1

financial_helps = [
    ('HLP_NO_ASK_HUSBAND', 'I don’t need to ask money from husband/family for my needs', 'अपनी जरूरतों के लिए पति या परिवार से पैसे मांगने की जरूरत नहीं पड़ती', 'खुद रा खर्चा सारु धणी या घर में पीसा नीं मांगणा पड़ै'),
    ('HLP_BIGGEST_SOURCE', 'Income from enterprise is biggest source of income for family', 'उद्यम की आय परिवार की आय का सबसे बड़ा स्रोत है', 'धंधे री कमाई घर री कमाई रो सबसूं बड़ो साधन छै'),
    ('HLP_CHILD_EDUCATION', 'Income used in covering education expenses for children', 'बच्चों की पढ़ाई-लिखाई के खर्च उठाने में उपयोग हुई', 'टाबरां री पढ़ाई रो खारचो संभाल्यो'),
    ('HLP_PAY_DEBTS', 'I have been able to pay the family debts', 'परिवार का पुराना कर्ज चुकाने में सक्षम हुई', 'घर रो पुरानो कर्ज उतार्यो'),
    ('HLP_ACQUIRE_ASSETS', 'Contributed money in acquiring assets for family', 'परिवार के लिए परिसंपत्तियां (जमीन/मकान/सामान) खरीदने में योगदान दिया', 'घर सारु जमीन/सम्पत्ति बणाई'),
    ('HLP_MARRIAGE_EXP', 'Contributed money for marriage expenses', 'विवाह/शादी-ब्याह के खर्चों में सहयोग दिया', 'ब्याव-शादी में पीसा लगाया')
]
hlp_ids = [h[0] for h in financial_helps]
for hid, en, hi, raj in financial_helps:
    add_var(hid, 'AppVariables', 'FinancialHelpImpact', 'Option', 'Enum', en, '', 'Impact Option', en, title_hi=hi, title_raj=raj)
add_var('Q_C_23_00', 'Survey', 'FinancialHelpFromIncome', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'EnumList', 'How has income from enterprise helped you financially? (Multiselect)', 'Section C Q23', 'Financial impact', var_list=' , '.join(hlp_ids), title_hi='उद्यम से हुई आय ने आपकी आर्थिक रूप से कैसे मदद की?', title_raj='धंधे री कमाई सूं थारी आर्थिक हालत में कांई सुधार हुयो?', decimal=ord_idx)
ord_idx += 1

add_var('Q_C_23_01', 'Survey', 'FinancialHelp_EducationAmt', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', 'Education expenses contribution amount (Rs)', 'Section C Q23.1', 'Education amount', title_hi='शिक्षा में योगदान राशि (रु)', title_raj='पढ़ाई में कित्ता पीसा लगाया (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_23_02', 'Survey', 'FinancialHelp_DebtsAmt', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', 'Family debts paid amount (Rs)', 'Section C Q23.2', 'Debts paid amount', title_hi='कर्ज चुकाने में योगदान राशि (रु)', title_raj='कर्ज उतारण में कित्ता पीसा लगाया (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_23_03', 'Survey', 'FinancialHelp_AssetsAmt', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', 'Assets acquisition contribution amount (Rs)', 'Section C Q23.3', 'Assets amount', title_hi='परिसंपत्ति निर्माण में योगदान राशि (रु)', title_raj='सम्पत्ति लेवण में कित्ता पीसा लगाया (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_C_23_04', 'Survey', 'FinancialHelp_MarriageAmt', f'QuestionPrompt, SectionC, Order:{ord_idx}', 'Decimal', 'Marriage expenses contribution amount (Rs)', 'Section C Q23.4', 'Marriage amount', title_hi='विवाह खर्च में योगदान राशि (रु)', title_raj='ब्याव-शादी में कित्ता पीसा लगाया (रु)', decimal=ord_idx)
ord_idx += 1

# ==========================================
# 5. SECTION D: EASE OF DOING BUSINESS & CHALLENGES
# ==========================================
family_responses = [
    ('FAM_INDEPENDENT', 'I am able to run my enterprise independently', 'मैं अपना उद्यम पूरी तरह स्वतंत्र रूप से चलाने में सक्षम हूँ', 'म्हैं खुद अकेली धंधो चलाऊं'),
    ('FAM_NOT_IN_FAVOUR', 'My husband was not in favour of starting the business', 'शुरुआत में पति व्यवसाय शुरू करने के पक्ष में नहीं थे', 'पैली धणी राजी कोनी हा'),
    ('FAM_NOT_SUPPORTIVE_INITIAL', 'Husband not supportive initially, but now helps when required', 'पति पहले असहयोगी थे पर अब जरूरत पड़ने पर मदद करते हैं', 'पैली साथ कोनी देता पण अबै मदद करै'),
    ('FAM_SUPPORTED_SHG_LOAN', 'Since enterprise started with SHG loan, husband supported', 'क्योंकि काम SHG लोन से शुरू हुआ, पति ने पूरा समर्थन दिया', 'समूह सूं लोन मिल्यो तो धणी साथ दियो'),
    ('FAM_INITIAL_INVESTMENT', 'My husband/family supported me with initial investment', 'पति/परिवार ने शुरुआती पूंजी में पूरा सहयोग दिया', 'शुरुआती पीसा घर का दिया'),
    ('FAM_FULL_SUPPORT', 'I have full support of husband/family and helped in every way', 'पति व परिवार का हर संभव तरीके से पूरा समर्थन मिला', 'घर का रो पूरो सहयोग मिल्यो'),
    ('FAM_EFFECTIVE_WITH_SUPPORT', 'I run/could run enterprise more effectively with family support', 'परिवार के सहयोग से व्यवसाय और बेहतर चला पा रही हूँ', 'घर रा जणां रे सहयोग सूं काम बढ़िया चालै')
]
fam_ids = [f[0] for f in family_responses]
for fid, en, hi, raj in family_responses:
    add_var(fid, 'AppVariables', 'FamilySupportResponse', 'Option', 'Enum', en, '', 'Family Support Option', en, title_hi=hi, title_raj=raj)
add_var('Q_D_01_00', 'Survey', 'HusbandFamilyResponse', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'EnumList', 'How has been your husband’s/Family’s response towards your enterprise? (Multiselect)', 'Section D Q1', 'Family support level', var_list=' , '.join(fam_ids), title_hi='आपके उद्यम के प्रति आपके पति/परिवार का क्या रुख रहा है?', title_raj='थारे काम-धंधे सारु थारे धणी/कुटुंब रो कांई साथ रह्यो?', decimal=ord_idx)
ord_idx += 1

sourcing_comforts = [
    ('SRC_TRAVEL_ALONE', 'I travel alone and I handle negotiations independently', 'मैं अकेले यात्रा करती हूँ और मोलभाव खुद स्वतंत्र रूप से करती हूँ', 'अकेली जा र सौदो कर ल्याऊं'),
    ('SRC_TRAVEL_COMPANION', 'I need travel companion but I handle negotiations independently', 'साथी चाहिए पर मोलभाव खुद स्वतंत्र रूप से करती हूँ', 'साथी चाहिजे पण मोलतोल खुद करूं'),
    ('SRC_FAMILY_HANDLES', 'My family member handles the purchase', 'परिवार का सदस्य सामग्री खरीद का काम संभालता है', 'घर रो जणो माल लावै'),
    ('SRC_CRP_HELPS', 'OSF/SVEP CRP helps in sourcing material', 'OSF/SVEP सीआरपी सामग्री जुटाने में मदद करती हैं', 'सीआरपी दीदी माल मंगावण में मदद करै'),
    ('SRC_WANT_DIFF_NEED_SUPP', 'Want to source from different places but need support', 'अलग-अलग जगहों से माल लाना चाहती हूँ पर सहयोग चाहिए', 'दूजी जगां सूं माल लावणो चाहूं पण मदद चाहिजे'),
    ('SRC_CONTENT_NEARBY', 'I am content to source material from nearby market', 'पास के बाजार से ही सामग्री लेने में संतुष्ट हूँ', 'नेड़ले बाजार सूं माल लेवण में ही राजी हूँ')
]
src_comf_ids = [s[0] for s in sourcing_comforts]
for sid, en, hi, raj in sourcing_comforts:
    add_var(sid, 'AppVariables', 'SourcingComfort', 'Option', 'Enum', en, '', 'Comfort Option', en, title_hi=hi, title_raj=raj)
add_var('Q_D_02_00', 'Survey', 'MaterialSourcingComfort', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Enum', 'What is your level of comfort in sourcing material?', 'Section D Q2', 'Sourcing autonomy', var_list=' , '.join(src_comf_ids), title_hi='सामग्री खरीदने और लाने में आपका सहूलियत स्तर क्या है?', title_raj='बाजार सूं माल खरीदण में थानें कित्ती आसानी रहवै छै?', decimal=ord_idx)
ord_idx += 1

recoveries = [
    ('REC_NO_ISSUES', 'Yes, I don’t face any issues', 'हाँ, मुझे कोई समस्या नहीं आती', 'हाँ, कोई तकलीफ कोनी आवै'),
    ('REC_CASH_ONLY', 'Yes, but I conduct only cash transactions', 'हाँ, क्योंकि मैं केवल नकद लेन-देन ही करती हूँ', 'हाँ, क्यूंकि नकद ही काम करां'),
    ('REC_EVENTUALLY_PAYS', 'Yes, eventually everyone pays', 'हाँ, देर-सबेर सभी भुगतान कर देते हैं', 'हाँ, आज नीं तो काल सगळा दे देवै'),
    ('REC_LEARNT_NEGOTIATE', 'Yes, but I have learnt over the years how to negotiate', 'हाँ, समय के साथ मैंने तकादा करना सीख लिया है', 'हाँ, टेम रे सागे समझ आगी किकण तगादो करणो'),
    ('REC_HUSBAND_RECOVERS', 'No, but my husband is able to recover', 'नहीं, लेकिन मेरे पति वसूली कर लेते हैं', 'ना, पण धणी वसूली कर ल्यावै'),
    ('REC_LOSSES_DEBT', 'No, my business has suffered losses due to debt', 'नहीं, उधारी डूबने के कारण व्यवसाय को नुकसान हुआ है', 'ना, उधारी डूबबा सूं नुकसान हुयो')
]
rec_ids = [r[0] for r in recoveries]
for rid, en, hi, raj in recoveries:
    add_var(rid, 'AppVariables', 'DebtRecoveryStatus', 'Option', 'Enum', en, '', 'Recovery Option', en, title_hi=hi, title_raj=raj)
add_var('Q_D_03_00', 'Survey', 'CustomerPaymentRecovery', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Enum', 'Are you able to recover money from customers?', 'Section D Q3', 'Payment recovery', var_list=' , '.join(rec_ids), title_hi='क्या आप ग्राहकों से पैसे/उधारी की वसूली कर पाती हैं?', title_raj='कांई गिराहकां सूं पीसा/उधारी आराम सूं निकल जावै छै?', decimal=ord_idx)
ord_idx += 1

funding_experiences = [
    ('FND_SHG_SUFFICIENT', 'SHG loan is sufficient for current scale of business', 'व्यवसाय के वर्तमान पैमाने के लिए SHG लोन पर्याप्त है', 'अबार रा काम सारु समूह रो लोन घणो छै'),
    ('FND_PLOUGH_EARNINGS', 'I regularly plough in my business earnings', 'मैं अपने व्यवसाय की कमाई को दोबारा व्यवसाय में लगाती हूँ', 'कमाई पाछी धंधे में ही लगावां'),
    ('FND_SHG_SMALL', 'SHG loan size is smaller than my requirement', 'SHG ऋण की राशि मेरी जरूरत से कम है', 'समूह रो लोन म्हारी जरूरत सूं कम पड़ै'),
    ('FND_EASY_MFI', 'I get required loan easily from moneylender/NBFIs', 'साहूकार/NBFI से जरूरत के अनुसार आसानी से कर्ज मिल जाता है', 'साहूकार/कंपनी सूं लोन असानी सूं मिल जावै'),
    ('FND_FAMILY_HELPS', 'My family members help me with funds and loans', 'परिवार के सदस्य पूंजी और ऋण में मेरी मदद करते हैं', 'घर का पीसा री मदद करै'),
    ('FND_AVOID_HIGH_INT', 'Don’t prefer moneylender/NBFIs as interest rate is high', 'साहूकार/NBFI नहीं लेती क्योंकि ब्याज दर बहुत अधिक है', 'साहूकार/कंपनी रो ब्याज घणो लागै इस वास्ते नीं लेवां'),
    ('FND_AVOID_SHORT_TIME', 'Don’t prefer moneylender/NBFIs as repayment time is shorter', 'साहूकार/NBFI नहीं लेती क्योंकि चुकाने का समय बहुत कम होता है', 'चुकावण रो टेम घणो कम मिलै')
]
fnd_ids = [f[0] for f in funding_experiences]
for fid, en, hi, raj in funding_experiences:
    add_var(fid, 'AppVariables', 'FundingExperience', 'Option', 'Enum', en, '', 'Funding Option', en, title_hi=hi, title_raj=raj)
add_var('Q_D_04_00', 'Survey', 'FundingExperience', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'EnumList', 'What has been your experience in funding your business? (Multiselect)', 'Section D Q4', 'Funding experiences', var_list=' , '.join(fnd_ids), title_hi='व्यवसाय के लिए पूंजी जुटाने में आपका क्या अनुभव रहा है?', title_raj='पूंजी और लोन जुटावण में थारो कांई अनुभव रह्यो?', decimal=ord_idx)
ord_idx += 1

challenges = [
    ('CHL_OSF_PHASED', 'OSF phased out which affected fund sufficiency (Specify amount)', 'OSF समाप्त हो गया जिससे फंड की कमी हो गई (राशि लिखें)', 'OSF बंद होग्यो जिणसूं पीसा री तंगी होगी'),
    ('CHL_SCALE_UP_FUNDS', 'Need more funds to scale up business (Specify amount)', 'व्यवसाय बढ़ाने के लिए और पूंजी चाहिए (राशि लिखें)', 'काम बढ़ावण खातर और पीसा चाहिजे'),
    ('CHL_RENOVATE_FUNDS', 'Need large funds to renovate shop/premise (Specify amount)', 'दुकान की मरम्मत/सजावट हेतु बड़ी राशि चाहिए (राशि लिखें)', 'दुकान सुधारण खातर पीसा चाहिजे'),
    ('CHL_TIMELY_INPUTS', 'Need timely access to funds before production begins (Specify amount)', 'उत्पादन शुरू होने से पहले समय पर पूंजी चाहिए (राशि लिखें)', 'सीजन सूं पैली टेम पर पीसा चाहिजे'),
    ('CHL_ACCESS_MARKET', 'Need support to access bigger market at lower cost', 'कम लागत पर माल पाने हेतु बड़े बाजार तक पहुँच का समर्थन चाहिए', 'बड़ा बाजार सूं सस्ता में माल लेवण री मदद चाहिजे'),
    ('CHL_SELL_INVENTORY', 'Need help in selling my inventory (Specify current value)', 'बचा हुआ स्टॉक/इन्वेंटरी बेचने में मदद चाहिए (मूल्य लिखें)', 'पड़्यो माल बेचण में मदद चाहिजे'),
    ('CHL_LEARN_SOCIAL_MEDIA', 'Need help in learning use of social media', 'सोशल मीडिया का उपयोग सीखने में मदद चाहिए', 'मोबाइल/सोशल मीडिया सीखण में मदद चाहिजे'),
    ('CHL_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजी कोई अड़चन')
]
chl_ids = [c[0] for c in challenges]
for cid, en, hi, raj in challenges:
    add_var(cid, 'AppVariables', 'BusinessChallenge', 'Option', 'Enum', en, '', 'Challenge Option', en, title_hi=hi, title_raj=raj)
add_var('Q_D_05_00', 'Survey', 'CurrentChallenges', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'EnumList', 'What are the challenges you are facing now? (Multiselect)', 'Section D Q5', 'Active challenges', var_list=' , '.join(chl_ids), title_hi='वर्तमान में आप किन चुनौतियों का सामना कर रही हैं?', title_raj='अबार काम-धंधे में कांई-कांई अड़चनां आ री छै?', decimal=ord_idx)
ord_idx += 1

add_var('Q_D_05_01', 'Survey', 'Challenge_OSFPhasedOutAmt', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Decimal', 'OSF phased out fund deficit amount (Rs)', 'Section D Q5.1', 'OSF deficit amount', title_hi='OSF समाप्ति से हुई कमी की राशि (रु)', title_raj='OSF बंद होबा सूं कित्ता पीसा री कमी पड़ी (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_D_05_02', 'Survey', 'Challenge_ScaleUpFundAmt', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Decimal', 'Scale up funds required amount (Rs)', 'Section D Q5.2', 'Scale up amount', title_hi='व्यवसाय विस्तार हेतु आवश्यक राशि (रु)', title_raj='धंधो बढ़ावण सारु कित्ता पीसा चाहिजे (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_D_05_03', 'Survey', 'Challenge_RenovationFundAmt', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Decimal', 'Renovation funds required amount (Rs)', 'Section D Q5.3', 'Renovation amount', title_hi='दुकान सुधार हेतु आवश्यक राशि (रु)', title_raj='दुकान ठीक करावण सारु कित्ता पीसा चाहिजे (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_D_05_04', 'Survey', 'Challenge_TimelyInputsAmt', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Decimal', 'Timely inputs funds required amount (Rs)', 'Section D Q5.4', 'Inputs amount', title_hi='उत्पादन सामग्री हेतु समय पर आवश्यक राशि (रु)', title_raj='कच्चा माल सारु टेम पर कित्ता पीसा चाहिजे (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_D_05_05', 'Survey', 'Challenge_InventoryHelpAmt', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Decimal', 'Current value of unsold inventory (Rs)', 'Section D Q5.5', 'Inventory value', title_hi='बिना बिके स्टॉक का वर्तमान मूल्य (रु)', title_raj='पड़्या माल रो मोल कित्तो छै (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_D_05_06', 'Survey', 'Challenge_Other', f'QuestionPrompt, SectionD, Order:{ord_idx}', 'Text', 'Specify other challenge remark', 'Section D Q5.6', 'Other challenge remark', title_hi='अन्य चुनौती का विवरण दें', title_raj='दूजी अड़चन रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1

# ==========================================
# 6. SECTION E: IMPACT OF SVEP/OSF SCHEMES
# ==========================================
add_var('Q_E_01_00', 'Survey', 'AttendedTraining', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Enum', 'Have you attended any training under SVEP/OSF?', 'Section E Q1', 'Attended training', var_list='OPT_YES , OPT_NO', title_hi='क्या आपने SVEP/OSF के तहत कोई प्रशिक्षण लिया है?', title_raj='कांई थे SVEP/OSF रो कोई प्रशिक्षण लियो छो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_02_00', 'Survey', 'TrainingDetails', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Text', 'If Yes, specify training topic / trade', 'Section E Q2', 'Training topic', title_hi='यदि हाँ, तो प्रशिक्षण का विवरण दें', title_raj='अगर हाँ, तो कांई काम सीख्यो छो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_03_00', 'Survey', 'UsedTrainingComponent', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Enum', 'Did you use any training component in your enterprise?', 'Section E Q3', 'Used training', var_list='OPT_YES , OPT_NO', title_hi='क्या आपने उद्यम में प्रशिक्षण के किसी भाग का उपयोग किया?', title_raj='कांई सिख्योड़ो काम धंधे में काम आयो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_04_00', 'Survey', 'UsedTrainingDetails', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Text', 'If Yes, specify which component used', 'Section E Q4', 'Component used details', title_hi='यदि हाँ, तो बताएं क्या उपयोग किया', title_raj='अगर हाँ, तो कांई बात काम आई?', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_05_01', 'Survey', 'MonthlyIncomeBeforeLoan', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Decimal', 'Monthly income BEFORE the changes (Rs)', 'Section E Q5.1', 'Pre loan income', title_hi='ऋण से हुए बदलाव से पहले मासिक आय (रु)', title_raj='लोन लेवण सूं पैली महिना री कमाई (रु)', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_05_02', 'Survey', 'MonthlyIncomeAfterLoan', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Decimal', 'Monthly income AFTER the changes (Rs)', 'Section E Q5.2', 'Post loan income', title_hi='ऋण से हुए बदलाव के बाद मासिक आय (रु)', title_raj='लोन लेवण पाछै महिना री कमाई (रु)', decimal=ord_idx)
ord_idx += 1

crp_contributions = [
    ('CRP_SUBSIDY', 'Accessing subsidy', 'सब्सिडी प्राप्त करने में सहायता', 'सब्सिडी दिलावण में मदद'),
    ('CRP_DOCUMENTS', 'Getting necessary documents (Aadhar, PAN, Udyam, FSSAI)', 'आवश्यक दस्तावेज (आधार, पैन, उद्यम, FSSAI) बनवाना', 'जरूरी कागज (आधार, पैन, उद्यम) बणावण में'),
    ('CRP_BIZ_PLANS', 'Helped us understand business plans', 'व्यावसायिक योजना बनाने और समझने में मदद की', 'व्यापार री योजना बणावण सिखायो'),
    ('CRP_FEEDBACK', 'Gave critical feedback on business idea to improve operations', 'व्यावसायिक विचार पर महत्वपूर्ण सुझाव दिए जिससे काम सुधरा', 'सुझाव दिया जिणसूं काम सुधर्यो'),
    ('CRP_RECORDS', 'Trained us on maintaining records which we didn’t know earlier', 'बही-खाता और लिखित रिकॉर्ड रखने का प्रशिक्षण दिया', 'हिसाब-किताब रखणो सिखायो'),
    ('CRP_NEW_IDEAS', 'Gave us new ideas to increase our income from enterprise', 'आय बढ़ाने के नए विचार और तरीके बताए', 'कमाई बढ़ावण रा नवा विचार दिया'),
    ('CRP_BANK_LOANS', 'Helped in accessing loans from bank', 'बैंक से ऋण प्राप्त करने में सहयोग किया', 'बैंक सूं लोन करवायो'),
    ('CRP_COMMUNICATION', 'Helped in our communication skills', 'बातचीत और आत्मविश्वास में सुधार कराया', 'बोलचाल और आत्मविश्वास बढ़ायो'),
    ('CRP_MARKETING', 'Helped in marketing', 'मार्केटिंग और प्रचार में मदद की', 'मार्केटिंग में मदद करी'),
    ('CRP_COMPETITORS', 'Helped understand competitors and suggested ways to beat competition', 'प्रतिस्पर्धा समझने और बेहतर सेवा देने के तरीके बताए', 'प्रतिस्पर्धा सूं निपटण रा तरीका बताया')
]
crp_con_ids = [c[0] for c in crp_contributions]
for cid, en, hi, raj in crp_contributions:
    add_var(cid, 'AppVariables', 'CRPContribution', 'Option', 'Enum', en, '', 'CRP Option', en, title_hi=hi, title_raj=raj)
add_var('Q_E_06_00', 'Survey', 'CRPContributions', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'EnumList', 'What has been the contribution of SVEP/OSF CRP in your enterprise? (Multiselect)', 'Section E Q6', 'CRP contributions', var_list=' , '.join(crp_con_ids), title_hi='आपके उद्यम में SVEP/OSF सीआरपी का क्या योगदान रहा?', title_raj='सीआरपी दीदी रो थारे धंधे में कांई सहयोग रह्यो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_06_01', 'Survey', 'CRPContributionDocDetails', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'Text', 'Specify documents made with CRP help', 'Section E Q6.1', 'Document details', title_hi='सीआरपी की मदद से बनवाए गए दस्तावेजों का विवरण', title_raj='बणवाया कागजां रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1
add_var('Q_E_07_00', 'Survey', 'ExpectationsFromScheme', f'QuestionPrompt, SectionE, Order:{ord_idx}', 'LongText', 'What are your expectations from the SVEP/OSF scheme?', 'Section E Q7', 'Scheme expectations', title_hi='SVEP/OSF योजना से आपकी क्या अपेक्षाएं हैं?', title_raj='योजना सूं थारी और कांई आशावां छै?', decimal=ord_idx)
ord_idx += 1

# ==========================================
# 7. SECTION F: ONLINE TRANSACTIONS & SOCIAL MEDIA
# ==========================================
phone_opts = [
    ('PHN_YES', 'Yes', 'हाँ, खुद का स्मार्टफोन है', 'हाँ, खुद रो स्मार्टफोन छै'),
    ('PHN_NO', 'No', 'नहीं, स्मार्टफोन नहीं है', 'ना, स्मार्टफोन कोनी'),
    ('PHN_ACCESS', 'No, but I have access to smart phone', 'नहीं, लेकिन परिवार के स्मार्टफोन तक पहुंच है', 'ना, पण घर रा फोन सूं काम चालै')
]
phn_ids = [p[0] for p in phone_opts]
for pid, en, hi, raj in phone_opts:
    add_var(pid, 'AppVariables', 'PhoneAccess', 'Option', 'Enum', en, '', 'Phone Option', en, title_hi=hi, title_raj=raj)
add_var('Q_F_01_00', 'Survey', 'SmartphoneOwnership', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'Enum', 'Do you own a smart phone?', 'Section F Q1', 'Smartphone access', var_list=' , '.join(phn_ids), title_hi='क्या आपके पास स्मार्टफोन है?', title_raj='कांई थारे कनै स्मार्टफोन छै?', decimal=ord_idx)
ord_idx += 1
add_var('Q_F_02_00', 'Survey', 'UseQRUPI', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'Enum', 'Do you use QR code/mobile banking for money transactions?', 'Section F Q2', 'QR / UPI usage', var_list='OPT_YES , OPT_NO', title_hi='क्या आप पैसों के लेन-देन के लिए क्यूआर कोड / मोबाइल बैंकिंग का उपयोग करती हैं?', title_raj='कांई थे ऑनलाइन (क्यूआर कोड/यूपीआई) पीसा लेवो या देवो छो?', decimal=ord_idx)
ord_idx += 1

qr_counts = [
    ('QRC_5_10', '5-10', 'प्रतिदिन 5-10 लेन-देन', 'रोज 5-10 बार'),
    ('QRC_10_20', '10-20', 'प्रतिदिन 10-20 लेन-देन', 'रोज 10-20 बार'),
    ('QRC_20_40', '20-40', 'प्रतिदिन 20-40 लेन-देन', 'रोज 20-40 बार'),
    ('QRC_GT_40', 'More than 40', 'प्रतिदिन 40 से अधिक', 'रोज 40 सूं बत्ती बार')
]
qrc_ids = [q[0] for q in qr_counts]
for qid, en, hi, raj in qr_counts:
    add_var(qid, 'AppVariables', 'QRCount', 'Option', 'Enum', en, '', 'QR Count Option', en, title_hi=hi, title_raj=raj)
add_var('Q_F_03_00', 'Survey', 'QRDailyTransactions', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'Enum', 'If yes, daily how many transactions done using QR / UPI?', 'Section F Q3', 'QR transaction count', var_list=' , '.join(qrc_ids), title_hi='यदि हाँ, तो प्रतिदिन कितने लेन-देन क्यूआर कोड/मोबाइल बैंकिंग से होते हैं?', title_raj='अगर हाँ, तो रोज कित्ती बार ऑनलाइन लेन-देन हुवै?', decimal=ord_idx)
ord_idx += 1

qr_reasons = [
    ('QRN_NO_PHONE', 'Since I don’t own smart phone, it is difficult to transact', 'स्मार्टफोन न होने के कारण लेन-देन कठिन है', 'स्मार्टफोन नीं होबा सूं लेन-देन कोनी हुवै'),
    ('QRN_FEW_CUSTOMERS', 'Not many customers use smart phone for payments', 'हमारे क्षेत्र में ग्राहक ऑनलाइन भुगतान का उपयोग नहीं करते', 'गिराहक ऑनलाइन पीसा कोनी देवै'),
    ('QRN_DONT_KNOW', 'I don’t know how to use and monitor transactions', 'मुझे क्यूआर कोड/मोबाइल बैंकिंग का उपयोग और जांच करना नहीं आता', 'म्हानें ऑनलाइन लेन-देन करणो कोनी आवै')
]
qr_rsn_ids = [r[0] for r in qr_reasons]
for rid, en, hi, raj in qr_reasons:
    add_var(rid, 'AppVariables', 'QRNonReason', 'Option', 'Enum', en, '', 'QR Non Option', en, title_hi=hi, title_raj=raj)
add_var('Q_F_04_00', 'Survey', 'QRNonUseReason', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'Enum', 'If no, reason for not using QR code / mobile banking', 'Section F Q4', 'Reason for no QR', var_list=' , '.join(qr_rsn_ids), title_hi='यदि नहीं, तो क्यूआर कोड/मोबाइल बैंकिंग का उपयोग न करने का कारण', title_raj='अगर ना, तो ऑनलाइन काम नीं लेवण रो कारण?', decimal=ord_idx)
ord_idx += 1

social_platforms = [
    ('SOC_WHATSAPP', 'Whatsapp', 'व्हाट्सएप', 'व्हाट्सएप'),
    ('SOC_INSTA', 'Instagram', 'इंस्टाग्राम', 'इंस्टाग्राम'),
    ('SOC_PINTEREST', 'Pinterest', 'पिंटरेस्ट', 'पिंटरेस्ट'),
    ('SOC_FB', 'Facebook', 'फेसबुक', 'फेसबुक'),
    ('SOC_SNAP', 'Snapchat', 'स्नैपचैट', 'स्नैपचैट'),
    ('SOC_NONE', 'Don’t use social media', 'सोशल मीडिया का उपयोग नहीं करतीं', 'सोशल मीडिया कोनी चलावां')
]
soc_plat_ids = [s[0] for s in social_platforms]
for sid, en, hi, raj in social_platforms:
    add_var(sid, 'AppVariables', 'SocialPlatform', 'Option', 'Enum', en, '', 'Platform Option', en, title_hi=hi, title_raj=raj)
add_var('Q_F_05_00', 'Survey', 'SocialPlatformsUsed', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'EnumList', 'Which social media platforms do you use for your business? (Multiselect)', 'Section F Q5', 'Platforms used', var_list=' , '.join(soc_plat_ids), title_hi='आप अपने व्यवसाय के लिए कौन-से सोशल मीडिया प्लेटफॉर्म का उपयोग करती हैं?', title_raj='धंधे सारु कौन-सा सोशल मीडिया चलावो छो?', decimal=ord_idx)
ord_idx += 1

social_usages = [
    ('USOC_TEXTS', 'Use texts to ask/share prices and book orders', 'कीमत पूछने/साझा करने और ऑर्डर बुक करने के लिए मैसेजिंग', 'मैसेज सूं भाव बतावण और आर्डर लेवण'),
    ('USOC_IMAGES_PROMO', 'Share images to promote business', 'व्यवसाय प्रचार के लिए उत्पादों के फोटो साझा करना', 'धंधे रो प्रचार सारु फोटो भेजण'),
    ('USOC_IMAGES_VENDOR', 'Share images to enquire about availability of products to vendors', 'सप्लायर से माल की उपलब्धता पूछने हेतु फोटो भेजना', 'सप्लायर नें फोटो भेज र माल पूछण'),
    ('USOC_IDEAS', 'Get new ideas and information about new products/services', 'नए उत्पादों/सेवाओं के नए विचार व डिजाइन देखना', 'नवा डिजाइन और माल री जानकारी लेवण'),
    ('USOC_NONE', 'Don’t use social media', 'सोशल मीडिया का उपयोग नहीं करतीं', 'सोशल मीडिया कोनी काम लेवां')
]
usoc_ids = [u[0] for u in social_usages]
for uid, en, hi, raj in social_usages:
    add_var(uid, 'AppVariables', 'SocialUsageMode', 'Option', 'Enum', en, '', 'Usage Option', en, title_hi=hi, title_raj=raj)
add_var('Q_F_06_00', 'Survey', 'SocialPlatformUsageMode', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'EnumList', 'How do you use these platforms in your business? (Multiselect)', 'Section F Q6', 'Platform usage mode', var_list=' , '.join(usoc_ids), title_hi='आप अपने व्यवसाय में इन प्लेटफॉर्म्स का उपयोग कैसे करती हैं?', title_raj='धंधे में सोशल मीडिया किकण काम आवै छै?', decimal=ord_idx)
ord_idx += 1

social_freqs = [
    ('FRQ_DAILY', 'Daily', 'प्रतिदिन', 'रोज'),
    ('FRQ_2_3_WEEK', 'Twice or thrice a week', 'सप्ताह में 2-3 बार', 'हफ्ते में 2-3 बार'),
    ('FRQ_4_5_MONTH', 'Four-five times a month', 'महीने में 4-5 बार', 'महिना में 4-5 बार'),
    ('FRQ_OCCASIONS', 'Only on occasions', 'केवल विशेष अवसरों / त्योहारों पर', 'कदे-कदाई त्योहार पर'),
    ('FRQ_NONE', 'Don’t use social media', 'सोशल मीडिया का उपयोग नहीं करतीं', 'काम कोनी लेवां')
]
frq_ids = [f[0] for f in social_freqs]
for fid, en, hi, raj in social_freqs:
    add_var(fid, 'AppVariables', 'SocialFrequency', 'Option', 'Enum', en, '', 'Freq Option', en, title_hi=hi, title_raj=raj)
add_var('Q_F_07_00', 'Survey', 'SocialMediaFrequency', f'QuestionPrompt, SectionF, Order:{ord_idx}', 'Enum', 'How often do you use social media for your business?', 'Section F Q7', 'Frequency of social media', var_list=' , '.join(frq_ids), title_hi='आप व्यवसाय के लिए सोशल मीडिया का उपयोग कितनी बार करती हैं?', title_raj='धंधे सारु सोशल मीडिया कित्ती बार काम लेवो छो?', decimal=ord_idx)
ord_idx += 1

# ==========================================
# 8. SECTION G: POST-EXIT OSF IN BARAN & RATANGARH
# ==========================================
add_var('Q_G_01_00', 'Survey', 'OSFInterventionYear', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'Number', 'In which year was the OSF intervention made? (Year)', 'Section G Q1', 'Intervention year', title_hi='OSF योजना का हस्तक्षेप किस वर्ष हुआ था?', title_raj='OSF रो काम किस साल सुरू हुयो छो?', decimal=ord_idx)
ord_idx += 1

biz_ops_statuses = [
    ('BOS_SALE_REDUCED', 'Yes but the sale has reduced', 'हाँ, लेकिन बिक्री कम हो गई है', 'हाँ, पण बिक्री घटगी'),
    ('BOS_SCALE_INCREASED', 'Yes but the scale has increased', 'हाँ, और व्यवसाय का दायरा/बिक्री बढ़ी है', 'हाँ, और काम बढ़ग्यो'),
    ('BOS_SCALE_SAME', 'Yes, but the scale has remained the same', 'हाँ, लेकिन दायरा पहले जैसा ही रहा है', 'हाँ, पण काम पैली जेड़ो ही छै'),
    ('BOS_CLOSED', 'No, enterprise closed', 'नहीं, व्यवसाय बंद हो गया है', 'ना, काम-धंधो बंद होग्यो')
]
bos_ids = [b[0] for b in biz_ops_statuses]
for bid, en, hi, raj in biz_ops_statuses:
    add_var(bid, 'AppVariables', 'BusinessOperationalStatus', 'Option', 'Enum', en, '', 'Status Option', en, title_hi=hi, title_raj=raj)
add_var('Q_G_02_00', 'Survey', 'BusinessOperationalStatus', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'Enum', 'Is your business still operational?', 'Section G Q2', 'Current operational status', var_list=' , '.join(bos_ids), title_hi='क्या आपका व्यवसाय अभी भी चालू है?', title_raj='कांई थारो काम-धंधो अबार भी चालू छै?', decimal=ord_idx)
ord_idx += 1
add_var('Q_G_02_01', 'Survey', 'BusinessClosureYear', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'Number', 'If closed, specify year when it was closed', 'Section G Q2.1', 'Closure year', title_hi='यदि बंद हुआ, तो किस वर्ष में बंद हुआ?', title_raj='अगर बंद हुयो, तो किस साल बंद हुयो?', decimal=ord_idx)
ord_idx += 1

closing_reasons = [
    ('CLR_NO_GUIDE', 'Sales reduced over the years as there was no one guiding us', 'मार्गदर्शन न मिलने से बिक्री कम होती गई', 'कोई समझावण वाळो कोनी हो जिणसूं बिक्री घटगी'),
    ('CLR_NO_CAPITAL', 'Needed more capital to source material but there was no source of loan', 'माल लाने के लिए पूंजी चाहिए थी लेकिन कोई ऋण नहीं मिला', 'माल लावण सारु पूंजी कोनी मिली'),
    ('CLR_BANK_REFUSED', 'Banks refused to give us loan', 'बैंकों ने ऋण देने से मना कर दिया', 'बैंक लोन देबा सूं मना कर दियो'),
    ('CLR_NO_NEW_CUSTOMERS', 'Unable to reach new customers', 'नए ग्राहकों तक पहुँचने में असमर्थ रहे', 'नवा गिराहक कोनी मिल्या'),
    ('CLR_COMPETITORS', 'New competitors in the market offering discounts', 'बाजार में नए प्रतिस्पर्धियों ने छूट देकर ग्राहक तोड़ लिए', 'नवा दुकानदार छूट दे र गिराहक तोड़ लिया'),
    ('CLR_OTHER', 'Any other, specify', 'अन्य कोई कारण, विवरण दें', 'दूजो कोई कारण'),
    ('CLR_DONT_KNOW', 'Don’t know', 'पता नहीं', 'ठा कोनी')
]
clr_ids = [c[0] for c in closing_reasons]
for cid, en, hi, raj in closing_reasons:
    add_var(cid, 'AppVariables', 'ClosingReason', 'Option', 'Enum', en, '', 'Reason Option', en, title_hi=hi, title_raj=raj)
add_var('Q_G_03_00', 'Survey', 'ScalingDownClosingReasons', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'EnumList', 'What are the reasons for scaling down the business/closing the business?', 'Section G Q3', 'Scale down / close reasons', var_list=' , '.join(clr_ids), title_hi='व्यवसाय घटने या बंद होने के क्या कारण रहे?', title_raj='धंधो मंदो होबा या बंद होबा रा कांई कारण हा?', decimal=ord_idx)
ord_idx += 1
add_var('Q_G_03_01', 'Survey', 'ScalingDownOtherReason', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'Text', 'Specify other closing reason', 'Section G Q3.1', 'Other close reason remark', title_hi='अन्य कारण का विवरण दें', title_raj='दूजे कारण रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1

support_options = [
    ('SUP_CONT_LOAN', 'Continued access to OSF loan', 'OSF ऋण की निरंतर उपलब्धता', 'OSF रो लोन चालतो रैहतो'),
    ('SUP_CONT_CRP', 'Continued support by OSF CRPs', 'OSF सीआरपी का निरंतर सहयोग', 'सीआरपी दीदी रो साथ रैहतो'),
    ('SUP_OTHER', 'Any other, specify', 'अन्य कोई, विवरण दें', 'दूजी कोई सहायता')
]
sup_ids = [s[0] for s in support_options]
for sid, en, hi, raj in support_options:
    add_var(sid, 'AppVariables', 'SupportNeeded', 'Option', 'Enum', en, '', 'Support Option', en, title_hi=hi, title_raj=raj)
add_var('Q_G_04_00', 'Survey', 'SupportNeededForSustenance', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'EnumList', 'What kind of support could have helped you to manage your business?', 'Section G Q4', 'Support needed', var_list=' , '.join(sup_ids), title_hi='किस प्रकार की सहायता मिलने पर आप अपना व्यवसाय संभाल पातीं?', title_raj='कांई मदद मिलती तो थारो धंधो चालतो रैहतो?', decimal=ord_idx)
ord_idx += 1
add_var('Q_G_04_01', 'Survey', 'SupportNeededOther', f'QuestionPrompt, SectionG, Order:{ord_idx}', 'Text', 'Specify other support needed', 'Section G Q4.1', 'Other support remark', title_hi='अन्य आवश्यक सहायता का विवरण दें', title_raj='दूजी मदद रो ब्योरो लिखो', decimal=ord_idx)
ord_idx += 1

# ==========================================
# EXPORT APPVARIABLES.CSV
# ==========================================
appvars_csv_path = os.path.join(DATA_DIR, 'AppVariables.csv')
with open(appvars_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=STANDARD_APPVAR_HEADERS)
    writer.writeheader()
    writer.writerows(app_vars)

print(f"✅ Generated AppVariables.csv ({len(app_vars)} rows) at {appvars_csv_path}")

# ==========================================
# 9. BUILD SINGLE FLAT 'SURVEY' TABLE (ONE FORM, ONE TABLE)
# ==========================================
# Extract all Survey physical columns in order:
survey_questions = [r for r in app_vars if r['Table'] == 'Survey' and r['Tags'].startswith('QuestionPrompt')]
survey_physical_cols = [
    'ID', 'Status', 'InvestigatorID', 'CreatedOn', 'Latitude', 'Longitude'
]
for q in survey_questions:
    col_name = q['Column']
    if col_name and col_name not in survey_physical_cols:
        survey_physical_cols.append(col_name)

survey_csv_path = os.path.join(DATA_DIR, 'Survey.csv')
sample_survey_row = {c: '' for c in survey_physical_cols}
sample_survey_row['ID'] = 'SURV_001'
sample_survey_row['Status'] = 'Draft'
sample_survey_row['InvestigatorID'] = 'surveyor1@ommnomi.in'
sample_survey_row['CreatedOn'] = NOW_STR
sample_survey_row['Latitude'] = '25.2134'
sample_survey_row['Longitude'] = '76.8456'
sample_survey_row['District'] = 'DIST_BARAN'
sample_survey_row['Block'] = 'BLK_CHHIPABAROD'
sample_survey_row['VillageGP'] = 'Khandari'
sample_survey_row['RespondentName'] = 'Sunita Bai'
sample_survey_row['ContactNumber'] = '9876543210'
sample_survey_row['SHGName'] = 'Radha SHG'
sample_survey_row['VOName'] = 'Ujala VO'
sample_survey_row['CLFName'] = 'Pragati CLF'
sample_survey_row['SHGMembershipYears'] = '5'
sample_survey_row['LeadershipRole'] = 'OPT_YES'
sample_survey_row['LeadershipYears'] = '3'
sample_survey_row['RelatedToCRP'] = 'OPT_NO'
sample_survey_row['EPInterventionType'] = 'INT_OSF'
sample_survey_row['EnterpriseName'] = 'Sunita Kirana & General Store'
sample_survey_row['EnterpriseSetupYear'] = '2020'
sample_survey_row['LoanReceivedYear'] = '2021'
sample_survey_row['BusinessType'] = 'BTY_TRADING'
sample_survey_row['BusinessActivities'] = 'ACT_GROCERY , ACT_FANCY_STORE'
sample_survey_row['RespondentAge'] = 'AGE_26_35'
sample_survey_row['MaritalStatus'] = 'MAR_MARRIED'
sample_survey_row['SocialCategory'] = 'CST_OBC'
sample_survey_row['EducationStatus'] = 'EDU_8TH'
sample_survey_row['FamilyMemberCount'] = '5'
sample_survey_row['FamilyAdultsCount'] = '3'
sample_survey_row['FamilyChildrenCount'] = '2'
sample_survey_row['FamilyTotalEarning'] = '2'
sample_survey_row['FamilyMaleEarning'] = '1'
sample_survey_row['FamilyFemaleEarning'] = '1'
sample_survey_row['FamilyDisabledCount'] = '0'
sample_survey_row['FamilyIncomeSources'] = 'INC_AGRI , INC_RESP_ENT'
sample_survey_row['AnnualHouseholdIncome'] = 'INC_160K_200K'
sample_survey_row['ReasonsStartingBusiness'] = 'RSN_RISING_EXP , RSN_CRP_ENCOURAGED'
sample_survey_row['BusinessCycle'] = 'CYC_REGULAR_HOURS'
sample_survey_row['BusinessPlaceType'] = 'PLC_OWN'
sample_survey_row['LocationConvenience'] = 'LOC_CONVENIENT'
sample_survey_row['AnnualSalaryBill'] = 'SAL_NOT_REL'
sample_survey_row['MarketingMethods'] = 'MKT_SHOP_ONLY , MKT_SHG_MEETINGS'
sample_survey_row['SeasonalSalesMethod'] = 'SEA_NOT_REL'
sample_survey_row['SocialMediaForMarketing'] = 'SMM_WHATSAPP_ORDERS'
sample_survey_row['RecordKeepingHabit'] = 'RKH_AFTER_CRP_TRAIN'
sample_survey_row['RecordKeepingMethod'] = 'RKT_CRP_DIARY'
sample_survey_row['InitialStartCapital'] = '35000'
sample_survey_row['InitialCapitalArranged'] = 'CAP_SHG_LOAN'
sample_survey_row['SHGAssociationAssistance'] = 'SHG_INITIATE_LOAN , SHG_CRP_GUIDED'
sample_survey_row['MonthlyIncomeIncreaseByOSFSVEP'] = 'INC_3K_4K'
sample_survey_row['FinancialHelpFromIncome'] = 'HLP_NO_ASK_HUSBAND , HLP_CHILD_EDUCATION'
sample_survey_row['HusbandFamilyResponse'] = 'FAM_FULL_SUPPORT'
sample_survey_row['MaterialSourcingComfort'] = 'SRC_TRAVEL_ALONE'
sample_survey_row['CustomerPaymentRecovery'] = 'REC_CASH_ONLY'
sample_survey_row['FundingExperience'] = 'FND_SHG_SUFFICIENT'
sample_survey_row['AttendedTraining'] = 'OPT_YES'
sample_survey_row['TrainingDetails'] = 'Basic business management and inventory recording'
sample_survey_row['UsedTrainingComponent'] = 'OPT_YES'
sample_survey_row['UsedTrainingDetails'] = 'Daily recording in bahi-khata and stock checking'
sample_survey_row['MonthlyIncomeBeforeLoan'] = '4000'
sample_survey_row['MonthlyIncomeAfterLoan'] = '8000'
sample_survey_row['CRPContributions'] = 'CRP_BIZ_PLANS , CRP_RECORDS , CRP_BANK_LOANS'
sample_survey_row['SmartphoneOwnership'] = 'PHN_YES'
sample_survey_row['UseQRUPI'] = 'OPT_YES'
sample_survey_row['QRDailyTransactions'] = 'QRC_5_10'
sample_survey_row['SocialPlatformsUsed'] = 'SOC_WHATSAPP'
sample_survey_row['SocialPlatformUsageMode'] = 'USOC_TEXTS , USOC_IMAGES_PROMO'
sample_survey_row['SocialMediaFrequency'] = 'FRQ_DAILY'
sample_survey_row['OSFInterventionYear'] = '2021'
sample_survey_row['BusinessOperationalStatus'] = 'BOS_SCALE_INCREASED'
sample_survey_row['Status_Profile'] = 'SEC_DONE'
sample_survey_row['Status_Operations'] = 'SEC_DONE'
sample_survey_row['Status_Challenges'] = 'SEC_DONE'
sample_survey_row['Status_SchemeImpact'] = 'SEC_DONE'
sample_survey_row['Status_Digital'] = 'SEC_DONE'
sample_survey_row['Status_PostExit'] = 'SEC_DONE'


with open(survey_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=survey_physical_cols)
    writer.writeheader()
    writer.writerow(sample_survey_row)

print(f"✅ Generated Survey.csv ({len(survey_physical_cols)} physical columns) at {survey_csv_path}")

# ==========================================
# 10. BUILD APPUSER TABLE (SOP §3: ID, LASTEDITBY, LASTEDITON, DEV ACCOUNTS)
# ==========================================
appuser_cols = ['ID', 'Name', 'Email', 'Role', 'District', 'Language', 'DailyTarget', 'Status', 'LastEditBy', 'LastEditOn']
appusers = [
    {
        'ID': 'DevNomi',
        'Name': 'Nomeshwer Sharma',
        'Email': 'nomeshwer@ommnomi.in',
        'Role': 'ROLE_ADMIN',
        'District': 'DIST_JAIPUR',
        'Language': 'LANG_EN',
        'DailyTarget': '0',
        'Status': 'Active',
        'LastEditBy': 'DevNomi',
        'LastEditOn': NOW_STR
    },
    {
        'ID': 'DevHardi',
        'Name': 'Hardik Sharma',
        'Email': 'hardiksharma80912@gmail.com',
        'Role': 'ROLE_ADMIN',
        'District': 'DIST_JAIPUR',
        'Language': 'LANG_EN',
        'DailyTarget': '0',
        'Status': 'Active',
        'LastEditBy': 'DevHardi',
        'LastEditOn': NOW_STR
    },
    {
        'ID': 'USR_001',
        'Name': 'Kavita Meena',
        'Email': 'surveyor1@ommnomi.in',
        'Role': 'ROLE_INVESTIGATOR',
        'District': 'DIST_BARAN',
        'Language': 'LANG_HI',
        'DailyTarget': '10',
        'Status': 'Active',
        'LastEditBy': 'DevNomi',
        'LastEditOn': NOW_STR
    },
    {
        'ID': 'USR_002',
        'Name': 'Sunita Sharma',
        'Email': 'surveyor2@ommnomi.in',
        'Role': 'ROLE_INVESTIGATOR',
        'District': 'DIST_CHURU',
        'Language': 'LANG_RAJ',
        'DailyTarget': '10',
        'Status': 'Active',
        'LastEditBy': 'DevNomi',
        'LastEditOn': NOW_STR
    },
    {
        'ID': 'USR_003',
        'Name': 'Shehnaz Jahan',
        'Email': 'supervisor@ommnomi.in',
        'Role': 'ROLE_SUPERVISOR',
        'District': 'DIST_BARAN',
        'Language': 'LANG_HI',
        'DailyTarget': '20',
        'Status': 'Active',
        'LastEditBy': 'DevNomi',
        'LastEditOn': NOW_STR
    }
]
user_csv_path = os.path.join(DATA_DIR, 'AppUser.csv')
with open(user_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=appuser_cols)
    writer.writeheader()
    writer.writerows(appusers)

print(f"✅ Generated AppUser.csv ({len(appusers)} users with Language & DailyTarget) at {user_csv_path}")

# ==========================================
# 11. BUILD SAMPLINGFRAME TABLE
# ==========================================
sampling_cols = ['ID', 'District', 'Block', 'PilotZone', 'TargetQuota', 'sc_st_floor_pct', 'new_enterprise_floor_pct']
sampling_rows = [
    {'ID': 'SMP_001', 'District': 'DIST_BARAN', 'Block': 'BLK_CHHIPABAROD', 'PilotZone': 'Fertile East', 'TargetQuota': '30', 'sc_st_floor_pct': '30%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_002', 'District': 'DIST_BARAN', 'Block': 'BLK_BARAN', 'PilotZone': 'Fertile East', 'TargetQuota': '25', 'sc_st_floor_pct': '30%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_003', 'District': 'DIST_CHURU', 'Block': 'BLK_RATANGARH', 'PilotZone': 'Arid West', 'TargetQuota': '30', 'sc_st_floor_pct': '25%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_004', 'District': 'DIST_CHURU', 'Block': 'BLK_SUJANGARH', 'PilotZone': 'Arid West', 'TargetQuota': '25', 'sc_st_floor_pct': '25%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_005', 'District': 'DIST_DAUSA', 'Block': 'BLK_SIKANDRA', 'PilotZone': 'Semi-Arid', 'TargetQuota': '55', 'sc_st_floor_pct': '35%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_006', 'District': 'DIST_DUNGARPUR', 'Block': 'BLK_SAGWARA', 'PilotZone': 'Southern Tribal', 'TargetQuota': '30', 'sc_st_floor_pct': '70%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_007', 'District': 'DIST_DUNGARPUR', 'Block': 'BLK_GALIAKOT', 'PilotZone': 'Southern Tribal', 'TargetQuota': '25', 'sc_st_floor_pct': '70%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_008', 'District': 'DIST_JODHPUR', 'Block': 'BLK_MANDOR', 'PilotZone': 'Western Arid', 'TargetQuota': '20', 'sc_st_floor_pct': '20%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_009', 'District': 'DIST_JODHPUR', 'Block': 'BLK_LUNI', 'PilotZone': 'Western Arid', 'TargetQuota': '20', 'sc_st_floor_pct': '20%', 'new_enterprise_floor_pct': '50%'},
    {'ID': 'SMP_010', 'District': 'DIST_JODHPUR', 'Block': 'BLK_SHERGADH', 'PilotZone': 'Western Arid', 'TargetQuota': '15', 'sc_st_floor_pct': '20%', 'new_enterprise_floor_pct': '50%'}
]
sample_csv_path = os.path.join(DATA_DIR, 'SamplingFrame.csv')
with open(sample_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=sampling_cols)
    writer.writeheader()
    writer.writerows(sampling_rows)

print(f"✅ Generated SamplingFrame.csv ({len(sampling_rows)} quotas) at {sample_csv_path}")

# ==========================================
# 12. BUILD MASTER EXCEL WORKBOOKS (OMMNOMI BRANDED)
# ==========================================
OUTPUT_XLSX = os.path.join(PROJECT_DIR, "Universal_Dynamic_Survey_Engine.xlsx")
MASTER_XLSX = os.path.join(PROJECT_DIR, "CmF_SHG_Women_Entrepreneurs_Master_Database.xlsx")
DOWNLOADS_XLSX = r"C:\Users\hardi\Downloads\Universal_Dynamic_Survey_Engine.xlsx"

wb = openpyxl.Workbook()
wb.remove(wb.active)

header_font = Font(name='Roboto', size=11, bold=True, color='FFFFFF')
blue_fill = PatternFill(start_color='4285F4', end_color='4285F4', fill_type='solid') # #4285F4
green_fill = PatternFill(start_color='34A853', end_color='34A853', fill_type='solid') # #34A853
alt_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='DADCE0'), right=Side(style='thin', color='DADCE0'),
    top=Side(style='thin', color='DADCE0'), bottom=Side(style='thin', color='DADCE0')
)

csv_sheets = [
    ('Survey', 'Survey.csv', blue_fill),
    ('AppVariables', 'AppVariables.csv', blue_fill),
    ('AppUser', 'AppUser.csv', green_fill),
    ('SamplingFrame', 'SamplingFrame.csv', green_fill)
]

for sheet_title, csv_filename, fill_color in csv_sheets:
    csv_p = os.path.join(DATA_DIR, csv_filename)
    ws = wb.create_sheet(title=sheet_title)
    if os.path.exists(csv_p):
        with open(csv_p, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            for row_idx, row in enumerate(reader, 1):
                ws.append(row)
                for col_idx, cell_value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    if row_idx == 1:
                        cell.font = header_font
                        cell.fill = fill_color
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.font = Font(name='Roboto', size=10)
                        if row_idx % 2 == 0:
                            cell.fill = alt_fill

# Auto-fit column widths
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)

wb.save(OUTPUT_XLSX)
wb.save(MASTER_XLSX)
try:
    wb.save(DOWNLOADS_XLSX)
except Exception as e:
    print(f"Note: Could not save to downloads directly: {e}")

print(f"🎉 MASTER EXCEL WORKBOOKS CREATED:\n   - {OUTPUT_XLSX}\n   - {MASTER_XLSX}\n   - {DOWNLOADS_XLSX}")

# ==========================================
# 13. EXPORT COMPLETE QUESTION CATALOG (MD & CSV)
# ==========================================
catalog_md_path = os.path.join(PROJECT_DIR, "ALL_SURVEY_QUESTIONS.md")
catalog_csv_path = os.path.join(PROJECT_DIR, "ALL_SURVEY_QUESTIONS.csv")

catalog_headers = [
    '#', 'Column', 'QuestionID', 'Section', 'Type', 'English Prompt', 'Hindi Prompt', 'Rajasthani Prompt', 'Dropdown / Options'
]
catalog_rows = []

for idx, q in enumerate(survey_questions, 1):
    # Determine section from Tags
    sec = 'Section A: Basic Details'
    if 'SectionB' in q['Tags']:
        sec = 'Section B: Respondent & Household Profile'
    elif 'SectionC' in q['Tags']:
        sec = 'Section C: Enterprise Operations'
    elif 'SectionD' in q['Tags']:
        sec = 'Section D: Ease of Doing Business & Challenges'
    elif 'SectionE' in q['Tags']:
        sec = 'Section E: Impact of SVEP/OSF Schemes'
    elif 'SectionF' in q['Tags']:
        sec = 'Section F: Online Transactions & Social Media'
    elif 'SectionG' in q['Tags']:
        sec = 'Section G: Post-Exit OSF in Baran & Ratangarh'

    # Fetch options if any
    opts = ''
    if q['VariableList']:
        opt_ids = [x.strip() for x in q['VariableList'].split(',') if x.strip()]
        opt_objs = [o for o in app_vars if o['ID'] in opt_ids]
        opts = ' • ' + '<br>• '.join([f"{o['Title']} ({o['Title_hi']} / {o['Title_raj']})" for o in opt_objs])

    catalog_rows.append({
        '#': str(idx),
        'Column': q['Column'],
        'QuestionID': q['ID'],
        'Section': sec,
        'Type': q['ValueControl'],
        'English Prompt': q['Title'],
        'Hindi Prompt': q['Title_hi'],
        'Rajasthani Prompt': q['Title_raj'],
        'Dropdown / Options': opts
    })

# Write Catalog CSV
with open(catalog_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.DictWriter(f, fieldnames=catalog_headers)
    writer.writeheader()
    writer.writerows(catalog_rows)

# Write Catalog MD
with open(catalog_md_path, 'w', encoding='utf-8') as f:
    f.write("# Complete Trilingual Questionnaire Catalog (CINI 'One Form, One Table' Pattern)\n")
    f.write(f"> Extracted from: `Second draft Questionnaire for the Study on Performance of SHG - CMF Final Feedback.md`\n")
    f.write(f"> Total Survey Columns: **{len(catalog_rows)}** | Languages: **English, Hindi (हिन्दी), Rajasthani (राजस्थानी)**\n\n")
    f.write("| # | Column | Question ID | Section | Type | English Prompt | Hindi Prompt | Rajasthani Prompt | Dropdown Choices |\n")
    f.write("|---|---|---|---|---|---|---|---|---|\n")
    for r in catalog_rows:
        opts_formatted = r['Dropdown / Options'].replace('\n', ' ')
        f.write(f"| {r['#']} | `{r['Column']}` | `{r['QuestionID']}` | {r['Section']} | **{r['Type']}** | {r['English Prompt']} | {r['Hindi Prompt']} | {r['Rajasthani Prompt']} | {opts_formatted} |\n")

print(f"✅ Generated ALL_SURVEY_QUESTIONS.md and ALL_SURVEY_QUESTIONS.csv ({len(catalog_rows)} questions cataloged)")
