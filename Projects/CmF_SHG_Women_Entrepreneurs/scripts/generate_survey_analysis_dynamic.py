"""
OmmNoMi Dynamic Survey Analysis CLI Orchestrator.
Discovers data dynamically, builds Master Interactive Web Dashboard,
and generates multi-sheet Excel workbooks and CSV exports for all districts.
Strictly adheres to <= 300 lines per file policy.
"""

import os
import argparse
import csv
import pandas as pd
import openpyxl
from analysis_engine import (
    resolve_all_files,
    SchemaRegistry,
    build_finance_sheet,
    build_social_matrix_sheet,
    build_agency_sourcing_sheet,
    build_indicators_sheet,
    build_all_in_one_sheet,
    build_district_html_report,
    build_master_dashboard_html
)


def export_workbook_to_csv(wb: openpyxl.Workbook, csv_dir: str) -> None:
    """Exports all sheets from an openpyxl workbook into individual CSV files."""
    os.makedirs(csv_dir, exist_ok=True)
    for sheetname in wb.sheetnames:
        ws = wb[sheetname]
        clean_name = sheetname.replace('&', 'and').replace(' ', '_')
        clean_name = "".join(c for c in clean_name if c.isalnum() or c == '_')
        while '__' in clean_name:
            clean_name = clean_name.replace('__', '_')
        csv_file = os.path.join(csv_dir, f"{clean_name}.csv")
        with open(csv_file, 'w', newline='', encoding='utf-8') as f:
            writer = csv.writer(f)
            for row in ws.iter_rows(values_only=True):
                if any(v is not None for v in row):
                    writer.writerow([str(v) if v is not None else "" for v in row])


def export_all_in_one_workbooks(
    df_survey: pd.DataFrame,
    df_sub: pd.DataFrame,
    df_subsub: pd.DataFrame,
    schema,
    output_dir: str,
    cohort_name: str = "Dausa"
) -> dict:
    """Generates consolidated all-in-one workbooks (single stacked sheet and combined 5-tab master)."""
    os.makedirs(output_dir, exist_ok=True)

    # 1. Single Stacked Sheet Workbook (All 4 in 1 sheet)
    wb_single = openpyxl.Workbook()
    ws_single = wb_single.active
    ws_single.title = "All In One Analysis"
    build_all_in_one_sheet(ws_single, df_survey, df_sub, df_subsub, schema, cohort_name)
    p_single = os.path.join(output_dir, f"{cohort_name}_Survey_Analysis_Single_Sheet.xlsx")
    wb_single.save(p_single)

    # 2. Comprehensive Master Workbook (Stacked Sheet + 4 Individual Tabs)
    wb_master = openpyxl.Workbook()
    wb_master.remove(wb_master.active)
    
    ws_m_all = wb_master.create_sheet(title="All In One Master")
    build_all_in_one_sheet(ws_m_all, df_survey, df_sub, df_subsub, schema, cohort_name)
    
    ws_m_fin = wb_master.create_sheet(title="1. Finance & Capital")
    build_finance_sheet(ws_m_fin, df_survey, df_subsub, schema)
    
    ws_m_soc = wb_master.create_sheet(title="2. Social Category Matrix")
    build_social_matrix_sheet(ws_m_soc, df_survey, schema, cohort_name)
    
    ws_m_agc = wb_master.create_sheet(title="3. Agency & Sourcing")
    build_agency_sourcing_sheet(ws_m_agc, df_survey, schema, cohort_name)
    
    ws_m_ind = wb_master.create_sheet(title="4. Demographics & Indicators")
    build_indicators_sheet(ws_m_ind, df_survey, df_sub, df_subsub, schema, cohort_name)
    
    p_master = os.path.join(output_dir, f"{cohort_name}_Survey_Analysis_All_In_One_Master.xlsx")
    wb_master.save(p_master)

    return {'single_sheet': p_single, 'master_workbook': p_master}


def export_4_standalone_workbooks(
    df_survey: pd.DataFrame,
    df_sub: pd.DataFrame,
    df_subsub: pd.DataFrame,
    schema,
    output_dir: str,
    cohort_name: str = "Dausa"
) -> list:
    """Generates 4 separate standalone .xlsx files, one for each analytical sheet."""
    os.makedirs(output_dir, exist_ok=True)
    files = []

    # File 1: Finance & Capital Mobilization
    wb1 = openpyxl.Workbook()
    ws1 = wb1.active
    ws1.title = "Finance & Capital"
    build_finance_sheet(ws1, df_survey, df_subsub, schema)
    p1 = os.path.join(output_dir, "1_Finance_and_Capital_Mobilization.xlsx")
    wb1.save(p1)
    files.append(p1)

    # File 2: Social Category Matrix
    wb2 = openpyxl.Workbook()
    ws2 = wb2.active
    ws2.title = "Social Category Matrix"
    build_social_matrix_sheet(ws2, df_survey, schema, cohort_name)
    p2 = os.path.join(output_dir, "2_Social_Category_Matrix.xlsx")
    wb2.save(p2)
    files.append(p2)

    # File 3: Agency & Sourcing Independence
    wb3 = openpyxl.Workbook()
    ws3 = wb3.active
    ws3.title = "Agency & Sourcing"
    build_agency_sourcing_sheet(ws3, df_survey, schema, cohort_name)
    p3 = os.path.join(output_dir, "3_Agency_and_Sourcing_Independence.xlsx")
    wb3.save(p3)
    files.append(p3)

    # File 4: Demographics & Survey Indicators
    wb4 = openpyxl.Workbook()
    ws4 = wb4.active
    ws4.title = "Demographics & Indicators"
    build_indicators_sheet(ws4, df_survey, df_sub, df_subsub, schema, cohort_name)
    p4 = os.path.join(output_dir, "4_Demographics_and_Survey_Indicators.xlsx")
    wb4.save(p4)
    files.append(p4)

    return files


def run_district_analysis(
    district: str = "Dausa",
    data_dir: str = None,
    output_base: str = "projects/CmF_SHG_Women_Entrepreneurs/reports",
    custom_paths: dict = None,
    loaded_data: tuple = None
) -> dict:
    """Orchestrates analysis for a specified district (Excel, CSVs, and district HTML)."""
    dist_clean = district.strip().replace("DIST_", "").title()
    print(f"\n--- Processing District: {dist_clean} ---")

    if loaded_data:
        df_survey, df_sub, df_subsub, schema = loaded_data
    else:
        files = resolve_all_files(data_dir=data_dir, custom_paths=custom_paths)
        schema = SchemaRegistry(files['appvars'])
        df_raw = pd.read_csv(files['survey'])
        df_survey = df_raw[df_raw['Status'] == 'Working'].copy()
        if df_survey.empty and not df_raw.empty:
            df_survey = df_raw[df_raw['Status'] != 'Dummy'].copy()
        df_sub = pd.read_csv(files['sub'])
        df_subsub = pd.read_csv(files['subsub'])

    dist_filter = df_survey['District'].astype(str).str.contains(dist_clean, case=False, na=False)
    df_dist = df_survey[dist_filter].copy()
    n_resp = len(df_dist)
    if n_resp == 0:
        print(f"[WARN] No respondents found for district '{dist_clean}'. Skipping.")
        return {}

    sids = set(df_dist['ID'])
    df_dist_sub = df_sub[df_sub['Survey'].isin(sids)].copy()
    df_dist_subsub = df_subsub[df_subsub['Survey'].isin(sids)].copy()

    # Output Folders
    dist_dir = os.path.join(output_base, "district_analysis", dist_clean)
    excel_dir = os.path.join(dist_dir, "excel")
    html_dir = os.path.join(dist_dir, "html")
    csv_dir = os.path.join(dist_dir, "csv")
    for d in [excel_dir, html_dir, csv_dir]:
        os.makedirs(d, exist_ok=True)

    # Master Excel Workbook
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_fin = wb.create_sheet(title="Finance & Capital")
    build_finance_sheet(ws_fin, df_dist, df_dist_subsub, schema)

    ws_mat = wb.create_sheet(title="Social Category Matrix")
    build_social_matrix_sheet(ws_mat, df_dist, schema, dist_clean)

    ws_agc = wb.create_sheet(title="Agency & Sourcing")
    build_agency_sourcing_sheet(ws_agc, df_dist, schema, dist_clean)

    ws_ind = wb.create_sheet(title="Demographics & Indicators")
    build_indicators_sheet(ws_ind, df_dist, df_dist_sub, df_dist_subsub, schema, dist_clean)

    excel_path = os.path.join(excel_dir, f"{dist_clean}_Survey_Analysis_Master.xlsx")
    wb.save(excel_path)

    export_workbook_to_csv(wb, csv_dir)

    html_path = os.path.join(html_dir, f"{dist_clean}_Survey_Analysis_Report.html")
    build_district_html_report(df_dist, df_dist_sub, df_dist_subsub, schema, dist_clean, html_path)

    print(f"[OK] {dist_clean}: {n_resp} WE -> Excel, CSVs, and Report saved.")
    return {'district': dist_clean, 'respondents': n_resp, 'excel_path': excel_path, 'html_path': html_path}


def run_full_pipeline(
    target_district: str = "all",
    data_dir: str = None,
    output_base: str = "projects/CmF_SHG_Women_Entrepreneurs/reports",
    custom_paths: dict = None
) -> dict:
    """Executes the complete pipeline: Master Web Dashboard + District archives."""
    print("\n========================================================")
    print("  Starting OmmNoMi Dynamic Multi-District Pipeline")
    print("========================================================")

    files = resolve_all_files(data_dir=data_dir, custom_paths=custom_paths)
    schema = SchemaRegistry(files['appvars'])
    df_raw = pd.read_csv(files['survey'])
    # Clean and filter only Working status records
    df_survey = df_raw[df_raw['Status'] == 'Working'].copy()
    if df_survey.empty and not df_raw.empty:
        df_survey = df_raw[df_raw['Status'] != 'Dummy'].copy()
    df_sub = pd.read_csv(files['sub'])
    df_subsub = pd.read_csv(files['subsub'])
    loaded_data = (df_survey, df_sub, df_subsub, schema)

    # 1. Build Single Publishable Master Interactive HTML Dashboard
    master_dashboard_path = os.path.join(output_base, "CmF_Rajasthan_Survey_Analysis_Dashboard.html")
    build_master_dashboard_html(df_survey, df_sub, df_subsub, schema, master_dashboard_path)
    print(f"\n[OK] ⭐ Master Interactive Web Dashboard generated:\n     -> {master_dashboard_path} ({os.path.getsize(master_dashboard_path):,} bytes)")

    # 2. Build Consolidated State-Level Master Workbook & Archives
    all_dir = os.path.join(output_base, "district_analysis", "All_Rajasthan")
    for sub_d in ["excel", "csv"]:
        os.makedirs(os.path.join(all_dir, sub_d), exist_ok=True)
    wb_all = openpyxl.Workbook()
    wb_all.remove(wb_all.active)
    build_finance_sheet(wb_all.create_sheet(title="Finance & Capital"), df_survey, df_subsub, schema)
    build_social_matrix_sheet(wb_all.create_sheet(title="Social Category Matrix"), df_survey, schema, "All Rajasthan")
    build_agency_sourcing_sheet(wb_all.create_sheet(title="Agency & Sourcing"), df_survey, schema, "All Rajasthan")
    build_indicators_sheet(wb_all.create_sheet(title="Demographics & Indicators"), df_survey, df_sub, df_subsub, schema, "All Rajasthan")
    all_excel_path = os.path.join(all_dir, "excel", "All_Rajasthan_Survey_Analysis_Master.xlsx")
    wb_all.save(all_excel_path)
    export_workbook_to_csv(wb_all, os.path.join(all_dir, "csv"))
    print(f"[OK] All Rajasthan: {len(df_survey)} WE -> Consolidated State Master Workbook saved.")

    # 3. Build Standalone & All-in-One .xlsx Workbooks for Client Deliverables
    export_dir = os.path.join(output_base, "excel_exports")
    df_dausa = df_survey[df_survey['District'].astype(str).str.contains('DAUSA', case=False, na=False)].copy()
    sids_dausa = set(df_dausa['ID'])
    df_dausa_sub = df_sub[df_sub['Survey'].isin(sids_dausa)].copy()
    df_dausa_subsub = df_subsub[df_subsub['Survey'].isin(sids_dausa)].copy()
    
    export_4_standalone_workbooks(df_dausa, df_dausa_sub, df_dausa_subsub, schema, export_dir, "Dausa")
    export_all_in_one_workbooks(df_dausa, df_dausa_sub, df_dausa_subsub, schema, export_dir, "Dausa")
    print(f"[OK] 4 Standalone & All-In-One .xlsx Workbooks generated for Dausa (N={len(df_dausa)}) in: {export_dir}")

    # 4. Build District Archives
    raw_districts = df_survey['District'].dropna().unique()
    dist_results = {}
    if target_district.lower() == "all":
        for d_code in raw_districts:
            c_name = str(d_code).replace("DIST_", "").title()
            res = run_district_analysis(c_name, output_base=output_base, loaded_data=loaded_data)
            if res:
                dist_results[c_name] = res
    else:
        dist_results[target_district] = run_district_analysis(target_district, output_base=output_base, loaded_data=loaded_data)

    print("\n[OK] All district workbooks, CSV exports, and master dashboard complete.")
    return {'master_dashboard': master_dashboard_path, 'districts': dist_results, 'state_master': all_excel_path}


def main():
    parser = argparse.ArgumentParser(description="OmmNoMi Dynamic Multi-District Survey Analysis Engine")
    parser.add_argument("--district", default="all", help="Target district (default: 'all', or 'Dausa', 'Churu', etc.)")
    parser.add_argument("--data-dir", default=None, help="Directory containing survey CSV exports")
    parser.add_argument("--output-base", default="projects/CmF_SHG_Women_Entrepreneurs/reports", help="Base report directory")
    parser.add_argument("--survey", default=None, help="Explicit path to WCH - Survey.csv")
    parser.add_argument("--subtable", default=None, help="Explicit path to WCH - SubTable.csv")
    parser.add_argument("--subsubtable", default=None, help="Explicit path to WCH - SubSubTable.csv")
    parser.add_argument("--appvars", default=None, help="Explicit path to WCH - AppVariables.csv")

    args = parser.parse_args()
    custom_paths = {k: getattr(args, k) for k in ['survey', 'subtable', 'subsubtable', 'appvars'] if getattr(args, k)}

    run_full_pipeline(
        target_district=args.district,
        data_dir=args.data_dir,
        output_base=args.output_base,
        custom_paths=custom_paths if custom_paths else None
    )


if __name__ == "__main__":
    main()
