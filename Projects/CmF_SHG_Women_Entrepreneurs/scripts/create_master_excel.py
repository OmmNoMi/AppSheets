import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\data"
OUTPUT_XLSX = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\CmF_SHG_Women_Entrepreneurs_Master_Database.xlsx"
DOWNLOADS_XLSX = r"C:\Users\hardi\Downloads\CmF_SHG_Women_Entrepreneurs_Master_Database.xlsx"

wb = openpyxl.Workbook()
# remove default sheet
wb.remove(wb.active)

# Styling Definitions
header_font = Font(name='Roboto', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='4285F4', end_color='4285F4', fill_type='solid') # OmmNoMi Blue
alt_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='DADCE0'),
    right=Side(style='thin', color='DADCE0'),
    top=Side(style='thin', color='DADCE0'),
    bottom=Side(style='thin', color='DADCE0')
)

csv_sheets = [
    ('Survey', 'Survey.csv'),
    ('SurveyResponse', 'SurveyResponse.csv'),
    ('AppVariables', 'AppVariables.csv'),
    ('SamplingFrame', 'SamplingFrame.csv'),
    ('AppUser', 'AppUser.csv')
]

for sheet_title, csv_filename in csv_sheets:
    csv_path = os.path.join(DATA_DIR, csv_filename)
    ws = wb.create_sheet(title=sheet_title)
    
    if os.path.exists(csv_path):
        with open(csv_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            for row_idx, row in enumerate(reader, 1):
                ws.append(row)
                for col_idx, cell_value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    if row_idx == 1:
                        cell.font = header_font
                        cell.fill = header_fill
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.font = Font(name='Roboto', size=10)
                        if row_idx % 2 == 0:
                            cell.fill = alt_fill

# Add Questionnaire Guide Sheet
guide_csv = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\ALL_SURVEY_QUESTIONS.csv"
if os.path.exists(guide_csv):
    ws_guide = wb.create_sheet(title="Questionnaire_Guide")
    with open(guide_csv, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            ws_guide.append(row)
            for col_idx, cell_value in enumerate(row, 1):
                cell = ws_guide.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = PatternFill(start_color='34A853', end_color='34A853', fill_type='solid') # OmmNoMi Green
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = Font(name='Roboto', size=10)
                    if row_idx % 2 == 0:
                        cell.fill = alt_fill

# Auto-fit column widths across all sheets
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 60)

wb.save(OUTPUT_XLSX)
wb.save(DOWNLOADS_XLSX)
print(f"Master Excel Workbook saved to:\n  1. {OUTPUT_XLSX}\n  2. {DOWNLOADS_XLSX}")
