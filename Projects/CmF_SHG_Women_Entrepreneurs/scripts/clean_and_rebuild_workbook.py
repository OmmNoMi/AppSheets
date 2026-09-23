import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\data"
OUTPUT_XLSX = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\Universal_Dynamic_Survey_Engine.xlsx"
DOWNLOADS_XLSX = r"C:\Users\hardi\Downloads\Universal_Dynamic_Survey_Engine.xlsx"

# 1. Update SurveyResponse.csv to have ONLY clean physical columns
survey_resp_physical_cols = [
    'ID', 'Survey', 'Section', 'Question', 'AnswerOption', 'AnswerOptionList', 
    'AnswerNumber', 'AnswerText', 'QuestionOrder'
]
resp_csv_path = os.path.join(DATA_DIR, 'SurveyResponse.csv')
with open(resp_csv_path, 'w', newline='', encoding='utf-8-sig') as f:
    writer = csv.writer(f)
    writer.writerow(survey_resp_physical_cols)

print(f"Updated SurveyResponse.csv with {len(survey_resp_physical_cols)} physical columns.")

# 2. Rebuild the Clean Master Workbook
wb = openpyxl.Workbook()
wb.remove(wb.active)

header_font = Font(name='Roboto', size=11, bold=True, color='FFFFFF')
blue_fill = PatternFill(start_color='4285F4', end_color='4285F4', fill_type='solid')
green_fill = PatternFill(start_color='34A853', end_color='34A853', fill_type='solid')
alt_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='DADCE0'), right=Side(style='thin', color='DADCE0'),
    top=Side(style='thin', color='DADCE0'), bottom=Side(style='thin', color='DADCE0')
)

csv_sheets = [
    ('Survey', 'Survey.csv'),
    ('SurveyResponse', 'SurveyResponse.csv'),
    ('AppVariables', 'AppVariables.csv'),
    ('SamplingFrame', 'SamplingFrame.csv'),
    ('AppUser', 'AppUser.csv')
]

for sheet_title, csv_filename in csv_sheets:
    csv_path = os.path.join(DATA_DIR, csv_filename)
    ws = wb.create_sheet(title=sheet_title)
    
    if os.path.exists(csv_path):
        with open(csv_path, 'r', encoding='utf-8-sig') as f:
            reader = csv.reader(f)
            for row_idx, row in enumerate(reader, 1):
                ws.append(row)
                for col_idx, cell_value in enumerate(row, 1):
                    cell = ws.cell(row=row_idx, column=col_idx)
                    cell.border = thin_border
                    if row_idx == 1:
                        cell.font = header_font
                        cell.fill = blue_fill
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.font = Font(name='Roboto', size=10)
                        if row_idx % 2 == 0:
                            cell.fill = alt_fill

# Questionnaire Guide Sheet
guide_csv = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\ALL_SURVEY_QUESTIONS.csv"
if os.path.exists(guide_csv):
    ws_guide = wb.create_sheet(title="Questionnaire_Guide")
    with open(guide_csv, 'r', encoding='utf-8-sig') as f:
        reader = csv.reader(f)
        for row_idx, row in enumerate(reader, 1):
            ws_guide.append(row)
            for col_idx, cell_value in enumerate(row, 1):
                cell = ws_guide.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if row_idx == 1:
                    cell.font = header_font
                    cell.fill = green_fill
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.font = Font(name='Roboto', size=10)
                    if row_idx % 2 == 0:
                        cell.fill = alt_fill

# Column width auto-fit
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

wb.save(OUTPUT_XLSX)
wb.save(DOWNLOADS_XLSX)
print("Updated clean Master Workbook saved!")
