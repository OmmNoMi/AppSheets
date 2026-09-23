import os
import csv
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

DATA_DIR = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\data"
OUTPUT_XLSX = r"c:\Users\hardi\AppSheets\projects\CmF_SHG_Women_Entrepreneurs\Universal_Dynamic_Survey_Engine.xlsx"
DOWNLOADS_XLSX = r"C:\Users\hardi\Downloads\Universal_Dynamic_Survey_Engine.xlsx"

wb = openpyxl.Workbook()
wb.remove(wb.active)

header_font = Font(name='Roboto', size=11, bold=True, color='FFFFFF')
blue_fill = PatternFill(start_color='4285F4', end_color='4285F4', fill_type='solid') # Blue
green_fill = PatternFill(start_color='34A853', end_color='34A853', fill_type='solid') # Green
purple_fill = PatternFill(start_color='673AB7', end_color='673AB7', fill_type='solid') # Purple
alt_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='DADCE0'), right=Side(style='thin', color='DADCE0'),
    top=Side(style='thin', color='DADCE0'), bottom=Side(style='thin', color='DADCE0')
)

# 1. Survey Table (Parent)
ws_survey = wb.create_sheet(title='Survey')
survey_headers = [
    'ID', 'SurveyType', 'District', 'Block', 'VillageGP', 'CLFName', 'VOName', 
    'SHGName', 'RespondentName', 'EnterpriseName', 'ContactNumber', 'SetupYear', 
    'Status', 'InvestigatorID', 'CreatedOn', 'Latitude', 'Longitude'
]
ws_survey.append(survey_headers)
sample_survey = [
    'SURV_001', 'SHG_Women_Entrepreneurs', 'DIST_BARAN', 'BLK_CHHIPABAROD', 'Khandari', 
    'Pragati CLF', 'Ujala VO', 'Radha SHG', 'Sunita Devi', 'Sunita Kirana & General Store', 
    '9876543210', '2020', 'Draft', 'USR_001', '09/09/2026 10:30:00', '25.2134', '76.8456'
]
ws_survey.append(sample_survey)

# 2. SurveyResponse (Single AI-Ready Reporting Table)
ws_resp = wb.create_sheet(title='SurveyResponse')
resp_headers = [
    'ID', 'Survey', 'SurveyType', 'Section', 'QuestionID', 'QuestionText', 
    'QuestionText_hi', 'AnswerOption', 'AnswerOptionList', 'AnswerNumber', 
    'AnswerText', 'AnswerValue', 'QuestionOrder'
]
ws_resp.append(resp_headers)

# Read questions from AppVariables to pre-seed sample response structure
appvars_path = os.path.join(DATA_DIR, 'AppVariables.csv')
with open(appvars_path, 'r', encoding='utf-8-sig') as f:
    appvar_rows = list(csv.DictReader(f))

questions = [r for r in appvar_rows if r['ID'].startswith('Q_') and r['ID'] != 'Q_SECTIONS_LIST']

sec_names = {
    'SEC_B_PROFILE': 'Section B: Respondent & Household Profile',
    'SEC_C_OPERATIONS': 'Section C: Enterprise Operations & Capital',
    'SEC_D_CHALLENGES': 'Section D: Enterprise Challenges & Coping',
    'SEC_E_SCHEME_IMPACT': 'Section E: Impact of SVEP / OSF Schemes',
    'SEC_F_DIGITAL_MEDIA': 'Section F: Online Transactions & Social Media',
    'SEC_G_BARAN_EXIT': 'Section G: Status of Post-Exit OSF in Baran'
}

for idx, q in enumerate(questions, 1):
    sec = 'Section A: Basic Details'
    for skey, sval in sec_names.items():
        if skey in q['Tags']:
            sec = sval
            break
    
    ws_resp.append([
        f'RESP_001_{idx:03d}', 'SURV_001', 'SHG_Women_Entrepreneurs', sec, q['ID'],
        q['Title'], q['Title_hi'], '', '', '', '', '', idx
    ])

# 3. AppVariables (Dynamic Question & Choice Catalog)
ws_appvars = wb.create_sheet(title='AppVariables')
with open(appvars_path, 'r', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for row in reader:
        ws_appvars.append(row)

# 4. SamplingFrame
ws_sample = wb.create_sheet(title='SamplingFrame')
sample_csv = os.path.join(DATA_DIR, 'SamplingFrame.csv')
with open(sample_csv, 'r', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for row in reader:
        ws_sample.append(row)

# 5. AppUser
ws_user = wb.create_sheet(title='AppUser')
user_csv = os.path.join(DATA_DIR, 'AppUser.csv')
with open(user_csv, 'r', encoding='utf-8-sig') as f:
    reader = csv.reader(f)
    for row in reader:
        ws_user.append(row)

# 6. AI_Flat_Report_Template (One row per survey showing all answers)
ws_flat = wb.create_sheet(title='AI_Reporting_Template')
flat_headers = [
    'Survey_ID', 'Survey_Type', 'District', 'Block', 'Village', 'Respondent_Name', 
    'Enterprise_Name', 'Setup_Year', 'Status', 'Enumerator_Email', 'Survey_Date'
] + [q['Title'] for q in questions]
ws_flat.append(flat_headers)

# Apply Styling
for ws in [ws_survey, ws_resp, ws_appvars, ws_sample, ws_user, ws_flat]:
    for row_idx, row in enumerate(ws.iter_rows(), 1):
        for cell in row:
            cell.border = thin_border
            if row_idx == 1:
                cell.font = header_font
                cell.fill = purple_fill if ws.title == 'AI_Reporting_Template' else blue_fill
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.font = Font(name='Roboto', size=10)
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

# Column width auto-fit
for sheet in wb.worksheets:
    for col in sheet.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or '')
            if '\n' in val:
                val = max(val.split('\n'), key=len)
            max_len = max(max_len, len(val))
        sheet.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 50)

wb.save(OUTPUT_XLSX)
wb.save(DOWNLOADS_XLSX)
print("Universal Dynamic Survey Engine Workbook generated successfully!")
