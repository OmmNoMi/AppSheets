"""
Finance & Capital Sheet Builder for Survey Analysis Engine.
Generates Part A (Enterprise Activity x Capital Source Matrix)
and Part B (Loan Usage Purpose x Capital Source Matrix).
"""

from typing import Dict, List, Set, Any
import pandas as pd
from openpyxl.utils import get_column_letter
from .data_evaluator import matches_activity
from .excel_styler import (
    font_head, font_subhead, font_bold, font_regular,
    fill_header, fill_subhead, fill_total, fill_highlight,
    border_cell, border_total, align_center, align_left, align_right,
    FMT_CURRENCY, autofit_columns
)


def build_finance_sheet(
    ws,
    df_survey: pd.DataFrame,
    df_subsub: pd.DataFrame,
    schema
) -> None:
    """Builds the comprehensive Finance & Capital worksheet."""
    ws.views.sheetView[0].showGridLines = True
    activities = schema.get_business_activities()
    capital_sources = schema.get_capital_sources()
    loan_usages = schema.get_loan_usages()

    # Pre-index respondents per activity
    activity_map = {}
    for sector, num, title, codes in activities:
        matching = df_survey[df_survey['BusinessActivities'].apply(
            lambda x, c=codes: matches_activity(x, c)
        )]
        activity_map[num] = matching

    # Part A Top Headers (Row 1)
    for col_idx, text in [(4, "No. of enterprises"), (5, "Total funds invested"), (6, "Average funds invested")]:
        c = ws.cell(row=1, column=col_idx, value=text)
        c.font, c.fill, c.alignment = font_head, fill_header, align_center

    ws.merge_cells(start_row=1, start_column=7, end_row=1, end_column=20)
    c_loan = ws.cell(row=1, column=7, value="Total loan amount from different sources")
    c_loan.font, c_loan.fill, c_loan.alignment = font_head, fill_header, align_center

    c_perf = ws.cell(row=1, column=21, value="Total business performance observations for current year (No.)")
    c_perf.font, c_perf.fill, c_perf.alignment = font_head, fill_header, align_center

    # Subheaders (Row 2)
    subheaders = [
        (1, "Sector"), (2, "#"), (3, "Business activities"),
        (4, "Count"), (5, "Total (Rs)"), (6, "Average (Rs)")
    ] + [(i, name) for i, (name, _) in enumerate(capital_sources, start=7)] + [(21, "Current Year Obs (No.)")]

    for col_idx, text in subheaders:
        c = ws.cell(row=2, column=col_idx, value=text)
        c.font, c.fill, c.alignment, c.border = font_subhead, fill_subhead, align_center, border_cell

    # Part A Rows
    cur_r = 3
    sector_ranges: Dict[str, List[int]] = {}
    subtotal_rows: List[int] = []

    for i, (sector, num, title, codes) in enumerate(activities):
        if sector not in sector_ranges:
            sector_ranges[sector] = []
        r_matches = activity_map[num]
        n_ent = len(r_matches)
        sids = set(r_matches['ID'])

        source_sums: List[float] = []
        tot_invested = 0.0
        for src_name, src_qg in capital_sources:
            sub = df_subsub[(df_subsub['Survey'].isin(sids)) & (df_subsub['Question_Group'] == src_qg)]
            nums = sub[sub['Question'].isin([
                'SubSubTable_CapitalArranged_FirstYear',
                'SubSubTable_CapitalArranged_MidYear',
                'SubSubTable_CapitalArranged_ThisYear'
            ])]['Answer_Number'].dropna()
            val = float(nums.sum())
            source_sums.append(val)
            tot_invested += val

        avg_invested = (tot_invested / n_ent) if n_ent > 0 else 0.0

        perf_sub = df_subsub[(df_subsub['Survey'].isin(sids)) & (df_subsub['Question_Group'].isin([
            'SubTable_BusinessChanges_AvSales', 'SubTable_BusinessChanges_AvIncome',
            'SubTable_BusinessChanges_Value', 'SubTable_BusinessChanges_Assets'
        ]))]
        pos_perf = perf_sub[(perf_sub['Question'] == 'SubSubTable_BusinessChanges_CurrentYear') & (perf_sub['Answer_Number'] > 0)]
        n_perf = len(pos_perf['Survey'].unique())

        ws.cell(row=cur_r, column=1, value=sector if not sector_ranges[sector] else "").font = font_bold
        ws.cell(row=cur_r, column=2, value=num).font = font_regular
        ws.cell(row=cur_r, column=3, value=title).font = font_regular
        ws.cell(row=cur_r, column=4, value=n_ent).font = font_bold if n_ent > 0 else font_regular
        ws.cell(row=cur_r, column=5, value=tot_invested).font = font_bold if tot_invested > 0 else font_regular
        ws.cell(row=cur_r, column=6, value=avg_invested).font = font_regular

        ws.cell(row=cur_r, column=4).number_format = '#,##0'
        ws.cell(row=cur_r, column=5).number_format = FMT_CURRENCY
        ws.cell(row=cur_r, column=6).number_format = FMT_CURRENCY

        for c_idx, s_val in enumerate(source_sums, start=7):
            c_src = ws.cell(row=cur_r, column=c_idx, value=s_val)
            c_src.font = font_bold if s_val > 0 else font_regular
            c_src.number_format = FMT_CURRENCY

        c_p = ws.cell(row=cur_r, column=21, value=n_perf)
        c_p.font = font_bold if n_perf > 0 else font_regular
        c_p.number_format = '#,##0'

        for col_c in range(1, 22):
            ws.cell(row=cur_r, column=col_c).border = border_cell
            ws.cell(row=cur_r, column=col_c).alignment = align_center if col_c in [1, 2] else (align_left if col_c == 3 else align_right)

        sector_ranges[sector].append(cur_r)
        cur_r += 1

        is_last_in_sector = (i == len(activities) - 1) or (activities[i + 1][0] != sector)
        if is_last_in_sector:
            ws.cell(row=cur_r, column=1, value=f"{sector} TOTAL").font = font_bold
            ws.cell(row=cur_r, column=3, value=f"TOTAL {sector.upper()}").font = font_bold
            st_r, end_r = sector_ranges[sector][0], sector_ranges[sector][-1]

            ws.cell(row=cur_r, column=4, value=f"=SUM(D{st_r}:D{end_r})").font = font_bold
            ws.cell(row=cur_r, column=5, value=f"=SUM(E{st_r}:E{end_r})").font = font_bold
            ws.cell(row=cur_r, column=6, value=f"=IF(D{cur_r}>0, E{cur_r}/D{cur_r}, 0)").font = font_bold
            ws.cell(row=cur_r, column=4).number_format = '#,##0'
            ws.cell(row=cur_r, column=5).number_format = FMT_CURRENCY
            ws.cell(row=cur_r, column=6).number_format = FMT_CURRENCY

            for c_idx in range(7, 21):
                cl = get_column_letter(c_idx)
                c_cell = ws.cell(row=cur_r, column=c_idx, value=f"=SUM({cl}{st_r}:{cl}{end_r})")
                c_cell.font, c_cell.number_format = font_bold, FMT_CURRENCY

            c_pu = ws.cell(row=cur_r, column=21, value=f"=SUM(U{st_r}:U{end_r})")
            c_pu.font, c_pu.number_format = font_bold, '#,##0'

            for col_c in range(1, 22):
                ws.cell(row=cur_r, column=col_c).fill = fill_total
                ws.cell(row=cur_r, column=col_c).border = border_cell
            subtotal_rows.append(cur_r)
            cur_r += 1

    # Grand Total Row for Part A
    ws.cell(row=cur_r, column=1, value="GRAND TOTAL").font = font_bold
    ws.cell(row=cur_r, column=3, value="OVERALL ENTERPRISE TOTAL").font = font_bold
    for c_idx in range(4, 22):
        cl = get_column_letter(c_idx)
        sum_str = "+".join([f"{cl}{sr}" for sr in subtotal_rows])
        if c_idx == 6:
            c_gt = ws.cell(row=cur_r, column=c_idx, value=f"=IF(D{cur_r}>0, E{cur_r}/D{cur_r}, 0)")
        else:
            c_gt = ws.cell(row=cur_r, column=c_idx, value=f"={sum_str}")
        c_gt.font = font_bold
        c_gt.number_format = FMT_CURRENCY if c_idx in range(5, 21) else '#,##0'
        c_gt.alignment = align_right

    for col_c in range(1, 22):
        ws.cell(row=cur_r, column=col_c).fill = fill_highlight
        ws.cell(row=cur_r, column=col_c).border = border_total
    cur_r += 4

    # Part B Header
    ws.merge_cells(start_row=cur_r, start_column=3, end_row=cur_r, end_column=18)
    p_b = ws.cell(row=cur_r, column=3, value="Distribution of loan amount from different sources as per the usage")
    p_b.font, p_b.fill, p_b.alignment = font_head, fill_header, align_center
    cur_r += 1

    # Part B Subheaders
    for c_idx, text in [(2, "#"), (3, "Loan Usage Purpose")] + [(i, name) for i, (name, _) in enumerate(capital_sources, start=4)] + [(18, "Total loan (Rs)")]:
        c = ws.cell(row=cur_r, column=c_idx, value=text)
        c.font, c.fill, c.alignment, c.border = font_subhead, fill_subhead, (align_left if c_idx == 3 else align_center), border_cell
    cur_r += 1

    st_usage_row = cur_r
    for u_num, u_title, u_code in loan_usages:
        ws.cell(row=cur_r, column=2, value=u_num).font = font_regular
        ws.cell(row=cur_r, column=2).alignment, ws.cell(row=cur_r, column=2).border = align_center, border_cell

        ws.cell(row=cur_r, column=3, value=u_title).font = font_regular
        ws.cell(row=cur_r, column=3).alignment, ws.cell(row=cur_r, column=3).border = align_left, border_cell

        for c_idx, (src_name, src_qg) in enumerate(capital_sources, start=4):
            matches = df_subsub[(df_subsub['Question_Group'] == src_qg) &
                                (df_subsub['Question'] == 'SubTable_CapitalLoanUsage_LoanUsage') &
                                (df_subsub['Answer_Enum'] == u_code)]
            m_surveys = matches['Survey'].unique()
            loan_recs = df_subsub[(df_subsub['Survey'].isin(m_surveys)) & (df_subsub['Question_Group'] == src_qg) &
                                  (df_subsub['Question'].isin([
                                      'SubSubTable_CapitalArranged_FirstYear',
                                      'SubSubTable_CapitalArranged_MidYear',
                                      'SubSubTable_CapitalArranged_ThisYear'
                                  ]))]
            sum_u = float(loan_recs['Answer_Number'].dropna().sum())
            c_u = ws.cell(row=cur_r, column=c_idx, value=sum_u)
            c_u.font = font_bold if sum_u > 0 else font_regular
            c_u.number_format, c_u.alignment, c_u.border = FMT_CURRENCY, align_right, border_cell

        tot_u = ws.cell(row=cur_r, column=18, value=f"=SUM(D{cur_r}:Q{cur_r})")
        tot_u.font, tot_u.number_format, tot_u.alignment, tot_u.border = font_bold, FMT_CURRENCY, align_right, border_cell
        cur_r += 1

    end_usage_row = cur_r - 1
    # Total All Usages Row
    ws.cell(row=cur_r, column=3, value="TOTAL ALL USAGES").font = font_bold
    ws.cell(row=cur_r, column=3).alignment, ws.cell(row=cur_r, column=3).fill, ws.cell(row=cur_r, column=3).border = align_left, fill_highlight, border_total

    for c_idx in range(4, 19):
        cl = get_column_letter(c_idx)
        tot_c = ws.cell(row=cur_r, column=c_idx, value=f"=SUM({cl}{st_usage_row}:{cl}{end_usage_row})")
        tot_c.font, tot_c.fill, tot_c.number_format, tot_c.alignment, tot_c.border = font_bold, fill_highlight, FMT_CURRENCY, align_right, border_total

    autofit_columns(ws, min_width=12, padding=3, custom_widths={'C': 38, 'U': 25})
