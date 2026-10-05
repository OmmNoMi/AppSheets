"""
Excel Styling & Formatting Utilities for Survey Analysis Engine.
Provides brand-consistent fonts, fills, borders, alignments, and block renderers.
"""

from typing import List, Tuple, Optional, Dict
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Number Formats (Double-quoted 'Rs ' avoids Excel date UserWarnings)
FMT_CURRENCY = '"Rs " #,##0'
FMT_PERCENT = '0.0%'
FMT_INT = '#,##0'
FMT_DECIMAL = '0.0'

# Brand Fonts
font_title = Font(name="Calibri", size=13, bold=True, color="1A73E8")
font_head = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
font_subhead = Font(name="Calibri", size=9.5, bold=True, color="174EA6")
font_bold = Font(name="Calibri", size=9.5, bold=True)
font_regular = Font(name="Calibri", size=9.5)
font_italic = Font(name="Calibri", size=9, italic=True, color="5F6368")

# Pattern Fills
fill_header = PatternFill(start_color="1A73E8", end_color="1A73E8", fill_type="solid")
fill_subhead = PatternFill(start_color="E8F0FE", end_color="E8F0FE", fill_type="solid")
fill_total = PatternFill(start_color="F1F3F4", end_color="F1F3F4", fill_type="solid")
fill_highlight = PatternFill(start_color="CEEAD6", end_color="CEEAD6", fill_type="solid")

# Borders
thin_side = Side(border_style="thin", color="DADCE0")
border_cell = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
border_total = Border(
    left=thin_side, right=thin_side,
    top=Side(border_style="thin", color="5F6368"),
    bottom=Side(border_style="double", color="202124")
)

# Alignments
align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
align_left = Alignment(horizontal="left", vertical="center")
align_right = Alignment(horizontal="right", vertical="center")


def render_table_block(
    ws,
    start_row: int,
    tbl_num: str,
    tbl_title: str,
    items: List[Tuple[str, int]],
    total_n: int,
    is_multiselect: bool = False
) -> int:
    """
    Renders a standard 2-4 column demographic/analytical indicator table block.
    Returns the next available row index.
    """
    r = start_row

    # Table Header Row
    ws.cell(row=r, column=1, value=f"Table {tbl_num}").font = font_bold
    ws.cell(row=r, column=2, value=tbl_title).font = font_bold

    c_cnt = ws.cell(row=r, column=3, value="Number of WE")
    c_cnt.font = font_subhead
    c_cnt.fill = fill_subhead
    c_cnt.alignment = align_center

    c_pct = ws.cell(row=r, column=4, value="% of Total WE")
    c_pct.font = font_subhead
    c_pct.fill = fill_subhead
    c_pct.alignment = align_center
    r += 1

    first_data_row = r
    for label, count in items:
        pct = (count / total_n) if total_n > 0 else 0.0

        c_lbl = ws.cell(row=r, column=2, value=label)
        c_lbl.font = font_regular
        c_lbl.alignment = align_left

        c_val = ws.cell(row=r, column=3, value=count)
        c_val.font = font_bold if count > 0 else font_regular
        c_val.number_format = FMT_INT
        c_val.alignment = align_right

        c_pc = ws.cell(row=r, column=4, value=pct)
        c_pc.font = font_regular
        c_pc.number_format = FMT_PERCENT
        c_pc.alignment = align_right

        for c_idx in range(2, 5):
            ws.cell(row=r, column=c_idx).border = border_cell
        r += 1

    last_data_row = r - 1

    # Total Row
    tot_label = "TOTAL RESPONSES" if is_multiselect else "TOTAL"
    c_tot_lbl = ws.cell(row=r, column=2, value=tot_label)
    c_tot_lbl.font = font_bold
    c_tot_lbl.alignment = align_left

    c_tot_cnt = ws.cell(
        row=r, column=3,
        value=f"=SUM(C{first_data_row}:C{last_data_row})"
    )
    c_tot_cnt.font = font_bold
    c_tot_cnt.number_format = FMT_INT
    c_tot_cnt.alignment = align_right

    c_tot_pct = ws.cell(
        row=r, column=4,
        value=f"=SUM(D{first_data_row}:D{last_data_row})"
    )
    c_tot_pct.font = font_bold
    c_tot_pct.number_format = FMT_PERCENT
    c_tot_pct.alignment = align_right

    for c_idx in range(2, 5):
        ws.cell(row=r, column=c_idx).fill = fill_total
        ws.cell(row=r, column=c_idx).border = border_total

    r += 2
    return r


def autofit_columns(
    ws,
    min_width: int = 12,
    padding: int = 3,
    custom_widths: Optional[Dict[str, int]] = None
) -> None:
    """Autofits all columns on a worksheet with optional custom overrides."""
    custom = custom_widths or {}
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        if col_letter in custom:
            ws.column_dimensions[col_letter].width = custom[col_letter]
            continue
        max_len = max(len(str(cell.value or '')) for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + padding, min_width)
